from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from backend.domain.inventory_reports import InventoryReportBatch, InventoryReportLookup, InventoryReportRow


class InventoryReportTests(APITestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser('report-test', password='test')
        self.client.force_authenticate(self.user)
        self.batch = InventoryReportBatch.objects.create(active=True, manifest={'date_from': '2026-01-01', 'date_to': '2026-12-31'})
        InventoryReportLookup.objects.bulk_create([
            InventoryReportLookup(batch=self.batch, kind='material', code=code, source_order=i, name='物料')
            for i, code in enumerate(('ZZJ-00001', 'ZZJ-00002', 'ZZJ-00003'), 1)
        ] + [InventoryReportLookup(batch=self.batch, kind='part_type', code=code, source_order=i)
             for i, code in enumerate(('包材类', '电子类', '塑胶类', '五金类'), 1)])
        self.url = '/api/material-inventory-detail/'

    def record(self, number, **changes):
        values = dict(batch=self.batch, kind='inventory', source_table='ld_det', source_key=str(number),
                      site='1000', material_code='ZZJ-00001', material_name='样品', specification='10%_规格',
                      location_code='ASM01', location_type='ASM', part_type='电子类',
                      payload={'material_code': 'ZZJ-00001', 'quantity': '0', 'inventory_amount': None, 'material_name': '=1+1'})
        values.update(changes)
        return InventoryReportRow.objects.create(**values)

    def get(self, **params):
        return self.client.get(self.url, {'source': 'dgyzx1', **params})

    def test_page_is_bounded_in_sql_and_keeps_zero_stock(self):
        for i in range(205):
            self.record(i)
        with CaptureQueriesContext(connection) as queries:
            result = self.get(page=2, page_size=100)
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(result.data['count'], 205)
        self.assertEqual(len(result.data['results']), 100)
        self.assertTrue(any('LIMIT 100 OFFSET 100' in q['sql'] for q in queries))
        self.assertEqual(result.data['results'][0]['quantity'], '0')
        self.assertEqual(self.get(page=3).data['results'].__len__(), 5)
        self.assertEqual(self.get(page_size=999999).status_code, 400)

    def test_nonzero_stock_filter_pagination_and_export(self):
        for i, quantity in enumerate(('0', '0.00000000', '0E-8', '-0.00000000', '12', '-3', '0.00000001')):
            self.record(i, payload={'quantity': quantity})
        for source in ('coerp_dgyzx1', 'local'):
            self.batch.source = source
            self.batch.save()
            params = {'source': 'local' if source == 'local' else 'dgyzx1', 'nonzero_quantity': '1'}
            result = self.client.get(self.url, {**params, 'page_size': 2, 'page': 2})
            self.assertEqual(result.status_code, 200, result.data)
            self.assertEqual(result.data['count'], 3)
            self.assertEqual(len(result.data['results']), 1)
            exported = self.client.post(self.url + 'export/', {'filters': params}, format='json')
            self.assertEqual(exported.status_code, 200)
            book = load_workbook(BytesIO(b''.join(exported.streaming_content)), read_only=True)
            exported.close()
            sheet = book['明细']
            values = list(sheet.values)
            column = values[0].index('库存数量')
            self.assertEqual({Decimal(str(row[column])) for row in values[1:]}, {Decimal('12'), Decimal('-3'), Decimal('0.00000001')})
            book.close()
            self.assertEqual(self.client.get(self.url, {**params, 'nonzero_quantity': '0'}).data['count'], 7)

    def test_ranges_source_order_wip_and_literal_contains(self):
        self.record(1)
        self.record(2, material_code='ZZJ-00002', location_type='WIP')
        self.record(3, material_code='ZZJ-00003', part_type='五金类')
        self.assertEqual(self.get(material_from='zzj-00001', material_to='ZZJ-00002').data['count'], 1)
        self.assertEqual(self.get(material_from='ZZJ-00001', material_to='ZZJ-00002', include_wip=1).data['count'], 2)
        self.assertEqual(self.get(material_from='ZZJ-00002', material_to='ZZJ-00001').status_code, 400)
        self.assertEqual(self.get(description='%_').data['count'], 2)
        self.assertEqual(self.get(description='not-present').data['count'], 0)
        self.url = '/api/material-movement-detail/'
        InventoryReportRow.objects.update(kind='movement', movement_type='120')
        self.assertEqual(self.get(part_type_from='包材类', part_type_to='塑胶类').data['count'], 2)

    def test_movement_date_boundary_type_conflicts_and_batch_pinning(self):
        self.url = '/api/material-movement-detail/'
        self.record(1, kind='movement', source_table='tr_hist', document_date='2026-03-01',
                    posted_at=timezone.make_aware(datetime(2026, 3, 2, 23, 59, 59)), movement_type='120')
        self.record(2, kind='movement', source_table='tr_hist', document_date='2026-03-02', movement_type='501')
        self.assertEqual(self.get(posted_from='2026-03-02', posted_to='2026-03-02').data['count'], 1)
        self.assertEqual(self.get(movement_type='501').status_code, 400)
        self.assertEqual(self.get(movement_type='501', exclude_cost=0).data['count'], 1)
        self.assertEqual(self.get(document_from='2026-03-03', document_to='2026-03-01').status_code, 400)
        InventoryReportBatch.objects.filter(pk=self.batch.pk).update(active=False)
        InventoryReportBatch.objects.create(active=True)
        self.assertEqual(self.get().data['count'], 0)
        self.assertEqual(self.get(batch_id=self.batch.pk).data['count'], 1)

    def test_export_all_filtered_rows_and_excel_text_safety(self):
        for i in range(105):
            self.record(i)
        result = self.client.post(self.url + 'export/', {'filters': {'batch_id': self.batch.pk}}, format='json')
        self.assertEqual(result.status_code, 200)
        content = b''.join(result.streaming_content)
        result.close()
        book = load_workbook(BytesIO(content), read_only=True)
        rows = list(book['明细'].values)
        self.assertEqual(len(rows), 106)
        keys = {value: i for i, value in enumerate(rows[0])}
        self.assertEqual(rows[1][keys['物料名称']], "'=1+1")
        self.assertEqual(rows[1][keys['库存数量']], 0)
        self.assertIsNone(book['明细'].cell(row=2, column=keys['库存金额'] + 1).value)

    def test_permission_and_lookup_contract(self):
        response = self.client.get(self.url + 'lookup/', {'kind': 'material', 'q': '00002'})
        self.assertEqual(response.data['count'], 1)
        self.assertEqual(set(response.data['results'][0]), {'code', 'name', 'site', 'specification'})
        self.client.force_authenticate(get_user_model().objects.create_user('no-access'))
        self.assertEqual(self.get().status_code, 403)
        self.assertEqual(self.client.get(self.url + 'lookup/').status_code, 403)
        self.assertEqual(self.client.post(self.url + 'export/', {}, format='json').status_code, 403)
