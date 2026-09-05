from pathlib import Path
from tempfile import TemporaryDirectory

from django.core.management import call_command
from django.test import TestCase
from openpyxl import Workbook
from unittest.mock import patch

from backend.models import Location, Material, ProductCategory, TaxCode, Uom, UomCategory
from backend.material_codes import allocate_material_code
from backend.test_support import SourceDatabaseIsolatedMixin


class TaxCodeImportTests(SourceDatabaseIsolatedMixin, TestCase):
    def test_imports_tax_code_and_invoice_name(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "tax.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["税务编码", "税务名称", "开票名称"])
            sheet.append(["1001", "塑料制品", "塑料制品（开票）"])
            workbook.save(path)
            call_command("import_tax_codes", path)

        tax = TaxCode.objects.get(code="1001")
        self.assertEqual(tax.name, "塑料制品")
        self.assertEqual(tax.invoice_name, "塑料制品（开票）")

    def test_material_syncs_tax_names_and_uses_max_category_serial(self):
        tax = TaxCode.objects.create(code="1001", name="塑料制品", invoice_name="塑料制品（开票）")
        category = ProductCategory.objects.create(code="A", name="成品", code_prefix="A")
        uom_category = UomCategory.objects.create(code="COUNT", name="数量")
        uom = Uom.objects.create(code="JIAN", name="件", category=uom_category)
        location = Location.objects.create(code="FG01", name="成品仓")
        existing = Material.objects.create(code="A-0007", name="已有物料", category=category, uom=uom,
                                           default_location=location, tax_code="1001", tax_master=tax)
        created = Material.objects.create(name="新物料", category=category, uom=uom, default_location=location,
                                          tax_code="1001")
        self.assertEqual(created.code, "A-0008")
        self.assertEqual(created.tax_name, tax.name)
        self.assertEqual(created.invoice_name, tax.invoice_name)
        self.assertEqual(existing.tax_name, tax.name)

    def test_import_reads_headers_on_every_sheet(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "tax.xlsx"
            workbook = Workbook()
            first = workbook.active
            first.append(["说明", "税务编码", "税务名称", "开票名称"])
            first.append(["A", "2001", "第一类", ""])
            second = workbook.create_sheet("补充")
            second.append(["税务名称", "开票名称", "税务编码"])
            second.append(["第二类", "第二类开票", "2002"])
            workbook.save(path)
            call_command("import_tax_codes", path)
        self.assertEqual(TaxCode.objects.get(code="2001").name, "第一类")
        self.assertEqual(TaxCode.objects.get(code="2002").invoice_name, "第二类开票")

    def test_source_serial_is_authoritative_and_preserves_width(self):
        with patch("backend.material_codes._source_codes", return_value=["ZXX-000123", "ZXX-000124"]):
            self.assertEqual(allocate_material_code("ZXX", ["ZXX-999999"]), "ZXX-000125")
            self.assertEqual(allocate_material_code("ZXX", ["ZXX-000125"]), "ZXX-000126")

    def test_empty_source_prefix_does_not_use_stale_local_serial(self):
        with patch("backend.material_codes._source_codes", return_value=[]):
            self.assertEqual(allocate_material_code("NEW", ["NEW-0099"]), "NEW-0001")
