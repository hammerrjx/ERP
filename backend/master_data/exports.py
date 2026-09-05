from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from urllib.parse import quote

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from rest_framework import serializers

from backend.models import (
    Currency, CustomerMaterial, Department, Location, Material, Partner,
    ProductCategory, SupplierQuote, Uom,
)


@dataclass(frozen=True)
class ExportColumn:
    label: str
    source: object = None
    kind: str = "text"


@dataclass(frozen=True)
class ExportDefinition:
    program: str
    sheet_name: str
    columns: tuple
    select_related: tuple = ()
    prefetch_related: tuple = ()


def c(label, source=None, kind="text"):
    return ExportColumn(label, source, kind)


def value_at(obj, source):
    if source is None:
        return ""
    if callable(source):
        return source(obj)
    value = obj
    for part in source.split("."):
        value = getattr(value, part, None)
        if value is None:
            return ""
    return value


def status_is(value):
    return lambda obj: obj.status == value


def location_type(obj):
    return {
        "fg": "FG", "rm": "RAW", "wip": "WIP", "ng": "NG",
        "virtual": "VIRTUAL", "warehouse": "WAREHOUSE",
    }.get(obj.location_type, obj.location_type)


def supply_code(obj):
    return {"purchase": "P", "production": "M", "outsource": "S"}.get(obj.supply_method, obj.supply_method)


def quality_required(obj):
    return obj.quality_control == "required"


def material_conversion(usage, field):
    def getter(obj):
        conversion = next((item for item in obj.uom_conversions.all() if item.usage == usage), None)
        value = getattr(conversion, field, "") if conversion else ""
        return value.name if field == "business_uom" and value else value
    return getter


def contact(number, field):
    def getter(obj):
        contacts = sorted(obj.contacts.all(), key=lambda item: (not item.is_primary, item.id))
        return getattr(contacts[number - 1], field, "") if len(contacts) >= number else ""
    return getter


def contact_columns(count):
    columns = []
    for number in range(1, count + 1):
        columns.extend((
            c(f"联系人{number}", contact(number, "name")),
            c(f"职位{number}", contact(number, "position")),
            c(f"联系电话{number}", contact(number, "phone")),
            c(f"传真{number}", contact(number, "fax")),
            c(f"电子邮箱{number}", contact(number, "email")),
        ))
    return tuple(columns)


AUDIT_COLUMNS = (
    c("录入人", "created_by"), c("录入时间", "created_at", "datetime"),
    c("修改次数"), c("最后修改人", "updated_by"), c("最后修改时间", "updated_at", "datetime"),
    c("审核", status_is("approved"), "bool"), c("审核人", "approved_by"), c("审核时间", "approved_at", "datetime"),
)


DEPARTMENT = ExportDefinition("AMDPMTA1", "1.4 -- 部门(AMDPMTA1)", (
    c("部门代码", "code"), c("部门名称", "name"), c("负责人", "manager"), c("备注", "notes"),
    c("上级部门", "parent.code"),
    c("录入人", lambda obj: obj.source_created_by or obj.created_by),
    c("录入时间", lambda obj: obj.source_created_at or obj.created_at, "datetime"),
), ("parent",))

CURRENCY = ExportDefinition("AMCUMTA1", "2.1 -- 货币(AMCUMTA1)", (
    c("货币", "code"), c("名称", "name"), c("符号", "symbol"),
    c("录入人", "created_by"), c("录入时间", "created_at", "datetime"),
    c("小数位", "decimal_places", "number"),
))

LOCATION = ExportDefinition("AMLOMTA1", "2.2 -- 库位(AMLOMTA1)", (
    c("库位", "code"), c("名称", "name"), c("类型", location_type),
    c("参与MRP", "participate_mrp", "bool"), c("可以使用", "usable", "bool"),
    c("可以负数", "allow_negative", "bool"), c("禁用库位", "disabled_for_inventory", "bool"),
    c("检验不良退货暂存仓", "quarantine_return", "bool"), c("仓管员", "manager"),
    c("审核权限", "approval_permissions"), c("录入人", "created_by"), c("录入日期", "created_at", "datetime"),
    c("修改次数"), c("修改人", "updated_by"), c("修改时间", "updated_at", "datetime"),
    c("存货会计科目", "storage_account"), c("主营业务收入科目", "sales_income_account"),
    c("发出商品成本科目", "shipped_goods_account"), c("主营业务成本科目", "sales_cost_account"),
    c("生产成本/直接材料科目", "material_account"), c("委托加工物资科目", "outsource_material_account"),
    c("存货会计科目名称"), c("主营业务收入科目名称"), c("仓管员名称"),
    c("发出商品科目名称"), c("主营业务成本科目名称"), c("生产成本/直接材料科目名称"),
    c("委托加工物资科目名称"), c("录入人姓名"), c("修改人姓名"),
    c("站点", "source_site"), c("所属部门", "department"), c("库位属性", "location_attribute"),
))

PRODUCT_CATEGORY = ExportDefinition("AMPLMTA1", "2.3 -- 产品类(AMPLMTA1)", (
    c("产品类", "code"), c("描述", "name"), c("零件类型", "part_type"), c("产品组", "product_group"),
    c("默认库位", "default_location.code"), c("采购员", "buyer"), c("采制代码", supply_code),
    c("虚项", "phantom", "bool"), c("发料", "issue_material", "bool"), c("非库存物料", "non_stock_material", "bool"),
    c("主生产计划", "main_production_plan", "bool"), c("关键物料", "key_material", "bool"), c("卷发", "roll", "bool"),
    c("质检", quality_required, "bool"), c("批号控制", "batch_control", "bool"), c("批号规则", "batch_rule"),
    c("物料查看权限", "material_view_permissions"), c("物料编辑权限", "material_edit_permissions"),
    c("物料删除权限", "material_delete_permissions"), c("物料审核/反审核权限", "material_approval_permissions"),
    c("存货会计科目", "inventory_account"), c("主营业务收入科目", "business_income_account"),
    c("委托加工物资科目", "outsource_material_account"), c("生产成本/直接材料科目", "production_cost_account"),
    c("主营业务成本科目", "business_cost_account"), c("录入人", "created_by"), c("录入时间", "created_at", "datetime"),
    c("可使用权限", "print_permission"), c("物料编码规则", "material_code_rule"),
    c("存货会计科目名称"), c("主营业务收入科目名称"), c("委托加工物资科目名称"),
    c("生产成本/直接材料科目名称"), c("主营业务成本科目名称"), c("零件类型名称"), c("产品组名称"),
    c("buyer_name"), c("默认单位", "default_uom.name"),
), ("default_location", "default_uom"))

UOM = ExportDefinition("AMUNMTA1", "2.4 -- 计量单位(AMUNMTA1)", (
    c("单位", "code"), c("描述", "name"), c("最小包装数量", "min_pack_qty", "number"),
    c("最小发出数量", "min_ship_qty", "number"), c("录入人", "created_by"),
    c("录入时间", "created_at", "datetime"), c("单位类别", "category.name"), c("换算系数", "factor", "number"),
), ("category",))

MATERIAL = ExportDefinition("AMPTMTA1", "2.5 -- 物料信息(AMPTMTA1)", (
    c("物料编码", "code"), c("物料名称", "name"), c("英文名称", "english_name"),
    c("海关物料编码", "customs_code"), c("海关物料名称", "customs_name"), c("条码", "barcode"),
    c("规格型号", "specification"), c("单位", "uom.name"), c("产品类", "category.code"),
    c("加入日期", "join_date", "date"), c("零件类型", "part_type"), c("产品组", "product_group"),
    c("绘图号", "drawing_number"), c("图像文件", lambda obj: obj.image.name if obj.image else ""), c("版本号", "version"),
    c("使用中", "active", "bool"), c("毛重(克)", "gross_weight_g", "number"), c("净重(克)", "net_weight_g", "number"),
    c("每箱产品数", "products_per_carton", "number"), c("损耗率", "scrap_rate", "number"), c("箱唛", "carton_mark"),
    c("长", "length", "number"), c("宽", "width", "number"), c("高", "height", "number"), c("体积", "volume", "number"),
    c("工程单位", material_conversion("engineering", "business_uom")),
    c("工程单位转换率(分子)", material_conversion("engineering", "business_qty"), "number"),
    c("工程单位转换率(分母)", material_conversion("engineering", "stock_qty"), "number"),
    c("采购单位", material_conversion("purchase", "business_uom")),
    c("采购单位转换率(分子)", material_conversion("purchase", "business_qty"), "number"),
    c("采购单位转换率(分母)", material_conversion("purchase", "stock_qty"), "number"),
    c("销售单位", material_conversion("sales", "business_uom")),
    c("销售单位转换率(分子)", material_conversion("sales", "business_qty"), "number"),
    c("销售单位转换率(分母)", material_conversion("sales", "stock_qty"), "number"),
    c("生产领用单位", material_conversion("production", "business_uom")),
    c("生产领用单位转换率(分子)", material_conversion("production", "business_qty"), "number"),
    c("生产领用单位转换率(分母)", material_conversion("production", "stock_qty"), "number"),
    c("ABC分类", "abc_class"), c("默认库位", "default_location.code"), c("默认储存位置", "default_storage_position"),
    c("生产线", "production_line"), c("盘点周期(天)", "count_cycle_days", "number"), c("批号控制", "batch_control", "bool"),
    c("批号规则", "batch_rule"), c("仓管员", "warehouse_manager"), c("保质期", "shelf_life_days", "number"),
    c("计划订单", "plan_order", "bool"), c("订购策略", "order_strategy"), c("订购数量", "order_qty", "number"),
    c("订购期间", "order_period_days", "number"), c("安全库存数量", "safety_stock_qty", "number"),
    c("库存警戒数量", "stock_warning_qty", "number"), c("采购员", "buyer"), c("计划员", "planner"),
    c("非生产物料", "non_production", "bool"), c("最大库存数量", "max_stock_qty", "number"),
    c("供应商", "default_supplier.code"), c("采制代码", supply_code), c("生产提前期(天)", "production_lead_days", "number"),
    c("采购提前期(天)", "purchase_lead_days", "number"), c("质检", quality_required, "bool"),
    c("质检提前期(天)", "quality_lead_days", "number"), c("虚项", "phantom", "bool"),
    c("最小订购数量", "min_purchase_qty", "number"), c("最小包装数量", "min_pack_qty", "number"),
    c("最小发出数量", "min_issue_qty", "number"), c("卷发", "roll", "bool"),
    c("超入库比例%", "over_receipt_ratio", "number"), c("不限制领料数量", "issue_qty_unrestricted", "bool"),
    c("不受采购报价权限限制", "purchase_quote_unrestricted", "bool"),
    c("不受销售报价权限限制", "sales_quote_unrestricted", "bool"),
    c("外协余料仓库存不参与MRP", "outsource_surplus_excluded_from_mrp", "bool"),
    c("默认工序", "default_operation", "number"), c("排程分类", "scheduling_class"), c("低层代码", "low_level_code", "number"),
    c("审核工程师", "engineering_reviewer"), c("生产入库扣料模式", "receipt_consume_mode"),
    c("外协收货扣料模式", "outsource_receipt_consume_mode"), c("签审状态", "status"),
    *AUDIT_COLUMNS,
    c("仓管员名称"), c("采购员名称"), c("计划员名称"), c("供应商简称", "default_supplier.short_name"),
    c("工程师名称"), c("库位名称", "default_location.name"), c("产线名称"), c("产品类描述", "category.name"),
    c("排程分类描述"), c("税务编码", "tax_code"), c("税务名称", "tax_name"), c("开票名称", "invoice_name"),
), ("uom", "category", "default_location", "default_supplier"), ("uom_conversions__business_uom",))

CUSTOMER = ExportDefinition("SLCMMTA1", "3.1 -- 客户资料(SLCMMTA1)", (
    c("客户代码", "code"), c("客户简称", "short_name"), c("客户全称", "name"), c("发票地址", "invoice_address"),
    c("经营地址", "operating_address"), c("送货地址", "delivery_address"), c("公司网址", "website"),
    *contact_columns(4),
    c("备注", "notes"), c("类别", "category"), c("类型", "business_type"), c("地区", "region"), c("报价方式", "quote_method"),
    c("课税类别", "tax_calculation_method"), c("支付方式", "payment_method"), c("月结日", "monthly_close_day", "number"),
    c("常用币种", "currency.code"), c("信用冻结", lambda obj: not obj.credit_enabled, "bool"), c("折扣率%", "discount_rate", "number"),
    c("发票类型", "invoice_type"), c("增值税率%", "tax_rate", "number"), c("备品比例%", "backup_ratio", "number"),
    c("超订单送货比例%", "over_order_delivery_ratio", "number"), c("寄售", "consignment", "bool"), c("税号", "tax_number"),
    c("付款银行", "payment_bank"), c("银行帐号", "bank_account"), c("法人", "legal_person"),
    c("注册资金(万元)", "registered_capital", "number"), c("营业额(万元)", "annual_revenue", "number"),
    c("开业日期", "opening_date", "date"), c("员工人数", "employee_count", "number"), c("客户总公司", "parent_company"),
    c("信用控制方案", "credit_policy"), c("一级额度", "credit_limit_1", "number"), c("二级额度", "credit_limit_2", "number"),
    c("超期天数1", "overdue_days_1", "number"), c("超期天数2", "overdue_days_2", "number"),
    c("预收帐款会计科目", "deposit_account"), c("应收帐款会计科目", "receivable_account"),
    c("主营业务收入会计科目", "income_account"), c("主营业务成本会计科目", "cost_account"),
    c("售后收入科目", "after_sales_income_account"), c("售后服务方式", "after_sales_method"), c("启用备货单"),
    c("铜产地", "copper_origin"), c("铜币种", "copper_currency"), c("报价类型", "price_type"),
    c("含税价计算方式", "tax_calculation_method"), c("订单按合计数量考虑数量级单价", "round_order_quantity_for_tier_price", "bool"),
    *AUDIT_COLUMNS,
    c("作废", status_is("void"), "bool"), c("作废人"), c("作废时间", kind="datetime"), c("签核状态", "status"),
    c("业务员名称", "sales_person"), c("支付方式名称", "payment_terms"), c("信用控制策略名称", "credit_policy"),
    c("预收帐款科目名称"), c("应收帐款科目名称"), c("主营业务收入科目名称"), c("主营业务成本科目名称"),
    c("售后收入科目名称"), c("售后服务方式名称", "after_sales_method"), c("确认", "is_confirmed", "bool"),
    c("确认人", "confirmed_by"), c("确认时间", "confirmed_at", "datetime"),
), ("currency",), ("contacts",))

CUSTOMER_MATERIAL = ExportDefinition("SLCPMTA1", "3.3 -- 客户物料(SLCPMTA1)", (
    c("客户代码", "customer.code"), c("物料编码", "material.code"), c("客户物料编码", "customer_code"),
    c("客户物料名称", "customer_name"), c("单位", "customer_uom.name"),
    c("单位转换率(分子)", "customer_uom_rate_m", "number"), c("单位转换率(分母)", "customer_uom_rate_d", "number"),
    c("备注", "notes"), c("录入人", "created_by"), c("录入时间", "created_at", "datetime"),
    c("客户简称", "customer.short_name"), c("物料名称", "material.name"), c("库存单位", "material.uom.name"),
    c("物料规格", "material.specification"), c("客户规格", "customer_specification"), c("客户条码", "customer_barcode"),
    c("终端客户编码", "terminal_customer_code"), c("终端客户名称", "terminal_customer_name"),
    c("启用", "enabled", "bool"), c("修改人", "updated_by"), c("修改时间", "updated_at", "datetime"),
), ("customer", "material__uom", "customer_uom"))

SUPPLIER = ExportDefinition("PUVDMTA1", "6.1 -- 供应商资料(PUVDMTA1)", (
    c("供应商代码", "code"), c("全称", "name"), c("简称", "short_name"), c("地址", "address"),
    c("联系人1", contact(1, "name")), c("联系电话1", contact(1, "phone")), c("传真1", contact(1, "fax")), c("电子邮箱1", contact(1, "email")),
    c("联系人2", contact(2, "name")), c("联系电话2", contact(2, "phone")), c("传真2", contact(2, "fax")), c("电子邮箱2", contact(2, "email")),
    c("网址", "website"), c("备注", "notes"), c("采购员", "buyer"), c("常用币种", "currency.code"),
    c("增值税率%", "tax_rate", "number"), c("支付方式", "payment_method"), c("备品比例%", "backup_ratio", "number"),
    c("超收货比例%", "over_receipt_ratio", "number"), c("服务评分", "service_score", "number"), c("报价方式", "quote_method"),
    c("发票类型", "invoice_type"), c("课税类别", "tax_category"), c("类型", "supplier_type"),
    c("按报价结算", "quoted_settlement", "bool"), c("税号", "tax_number"), c("法人", "legal_person"),
    c("注册资金(万元)", "registered_capital", "number"), c("年营业额(万元)", "annual_revenue", "number"),
    c("开业日期", "opening_date", "date"), c("员工人数", "employee_count", "number"), c("卖料单价取值", "material_price_method"),
    c("卖料单价系数", "material_price_factor", "number"), c("暂停付款", "payment_hold", "bool"),
    c("应付帐款科目", "purchase_account"), c("暂估应付帐科目", "temporary_payable_account"), c("委托加工费用科目"),
    c("委托加工物资科目", "outsource_material_account"), c("签核状态", "status"),
    *AUDIT_COLUMNS,
    c("作废", status_is("void"), "bool"), c("作废人"), c("作废时间", kind="datetime"), c("采购员名称", "buyer"),
    c("支付方式描述", "payment_terms"), c("应付帐款科目名称"), c("暂估应付科目名称"), c("委托加工费科目名称"),
    c("委托加工物资科目名称"), c("确认", "is_confirmed", "bool"), c("确认人", "confirmed_by"),
    c("确认时间", "confirmed_at", "datetime"), c("三证合一", "three_certificates", "bool"),
    c("开户许可证", "business_license", "bool"), c("类别", "category"), c("是否内部公司", "internal_company", "bool"),
), ("currency",), ("contacts",))


SUPPLIER_QUOTE = ExportDefinition("PUVQMTA1", "6.5 -- 供应商报价(PUVQMTA1)", (
    c("报价单号", "number"), c("供应商代码", "supplier.code"), c("币种", "currency.code"),
    c("业务类型", "business_group.code"),
    c("物料编码", "material.code"), c("采购单位", "purchase_uom.name"),
    c("单位转换率(分子)", "uom_rate_m", "number"), c("单位转换率(分母)", "uom_rate_d", "number"),
    c("生效日期", "effective_date", "date"), c("失效日期", "expiry_date", "date"),
    c("是否含税", "tax_included", "bool"), c("增值税率%", "tax_rate", "number"),
    c("报价类型", lambda obj: "P" if obj.quote_type == SupplierQuote.QuoteType.PURCHASE else "S"),
    c("材料单价", "material_unit_price", "number"), c("加工单价", "processing_unit_price", "number"),
    c("含税单价", "unit_price", "number"), c("最小采购数量", "min_purchase_qty", "number"),
    c("最小包装数量", "min_pack_qty", "number"), c("交货天数", "delivery_days", "number"),
    c("备注", "notes"), c("签核状态", "workflow_status"), c("录入人", "created_by"),
    c("录入时间", "created_at", "datetime"), c("修改次数", "modification_count", "number"),
    c("最后修改人", "updated_by"), c("最后修改时间", "updated_at", "datetime"),
    c("核准", "is_ratified", "bool"), c("核准人", "ratified_by"), c("核准时间", "ratified_at", "datetime"),
    c("物料名称", "material.name"), c("规格", "material.specification"), c("库存单位", "material.uom.name"),
    c("供应商简称", "supplier.short_name"), c("未税单价", "untaxed_unit_price", "number"),
    c("发票类型", "supplier.invoice_type"), c("审核", lambda obj: obj.status == "approved", "bool"),
    c("审核人", "approved_by"), c("审核时间", "approved_at", "datetime"),
    c("确认", "is_confirmed", "bool"), c("确认人", "confirmed_by"), c("确认时间", "confirmed_at", "datetime"),
    c("销售确认", "is_sales_confirmed", "bool"), c("销售确认人", "sales_confirmed_by"),
    c("采购确认时间", "sales_confirmed_at", "datetime"),
    c("工序", lambda obj: obj.operation.sequence if obj.operation_id else 0, "number"),
    c("主件编码", "parent_material.code"), c("税务编码", "material.tax_code"),
    c("税务名称", "material.tax_name"), c("开票名称", "material.invoice_name"), c("类别", "material.category.code"),
), ("supplier", "business_group", "currency", "material__uom", "material__category", "purchase_uom", "operation", "parent_material"), ("lines",))


EXPORT_DEFINITIONS = {
    Department: DEPARTMENT, Currency: CURRENCY, Location: LOCATION,
    ProductCategory: PRODUCT_CATEGORY, Uom: UOM, Material: MATERIAL,
    CustomerMaterial: CUSTOMER_MATERIAL,
}


def definition_for(model, partner_kind=None):
    if model is Partner:
        if partner_kind == "customer":
            return CUSTOMER
        if partner_kind == "supplier":
            return SUPPLIER
        raise serializers.ValidationError({"partner_kind": "客户或供应商导出必须指定资料类型"})
    return EXPORT_DEFINITIONS.get(model)


def export_queryset(queryset, request):
    definition = definition_for(queryset.model, request.data.get("partner_kind"))
    if definition is None:
        raise serializers.ValidationError("此资料暂不支持导出")
    ids_supplied = "ids" in request.data
    ids = request.data.get("ids")
    if ids_supplied:
        if not isinstance(ids, list) or any(not isinstance(value, int) or value <= 0 for value in ids):
            raise serializers.ValidationError({"ids": "当前筛选结果必须为正整数 ID 列表"})
        queryset = queryset.filter(pk__in=list(dict.fromkeys(ids)))
    partner_kind = request.data.get("partner_kind")
    if queryset.model is Partner:
        queryset = queryset.filter(kind__in=(partner_kind, "both"))
    if definition.select_related:
        queryset = queryset.select_related(*definition.select_related)
    if definition.prefetch_related:
        queryset = queryset.prefetch_related(*definition.prefetch_related)
    ordering = "code" if any(field.name == "code" for field in queryset.model._meta.fields) else "pk"
    return definition, queryset.order_by(ordering)


def safe_text(value):
    text = str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def excel_value(value, kind):
    if value is None or value == "":
        return ""
    if hasattr(value, "code") and not isinstance(value, (str, bytes)):
        value = value.code
    if kind == "bool":
        return bool(value)
    if kind == "number":
        return value if isinstance(value, (int, float, Decimal)) else Decimal(str(value))
    if kind in {"date", "datetime"}:
        if isinstance(value, datetime) and timezone.is_aware(value):
            value = timezone.localtime(value).replace(tzinfo=None)
        return value if isinstance(value, (date, datetime)) else ""
    return safe_text(value)


def build_workbook(definition, queryset):
    workbook = Workbook(write_only=True)
    workbook.properties.creator = "ERP"
    workbook.properties.title = definition.sheet_name
    sheet = workbook.create_sheet(definition.sheet_name)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for index, column in enumerate(definition.columns, 1):
        sheet.column_dimensions[_column_letter(index)].width = preferred_column_width(column)
    sheet.sheet_format.defaultRowHeight = 20
    sheet.row_dimensions[1].height = 28

    header_fill = PatternFill("solid", fgColor="DCE6F1")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="1F2937")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="center", wrap_text=False)
    grid_side = Side(style="thin", color="C7D0DA")
    header_border = Border(left=grid_side, right=grid_side, top=grid_side, bottom=grid_side)
    body_border = Border(left=grid_side, right=grid_side, top=grid_side, bottom=grid_side)
    stripe_fills = (
        PatternFill("solid", fgColor="FFFFFF"),
        PatternFill("solid", fgColor="F3F8FC"),
    )
    header = []
    for column in definition.columns:
        cell = WriteOnlyCell(sheet, value=column.label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = header_border
        header.append(cell)
    sheet.append(header)

    row_count = 1
    for record in queryset.iterator(chunk_size=500):
        data_row = row_count + 1
        sheet.row_dimensions[data_row].height = 20
        row = []
        for index, column in enumerate(definition.columns):
            value = excel_value(value_at(record, column.source), column.kind)
            cell = WriteOnlyCell(sheet, value=value)
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = body_alignment
            cell.border = body_border
            cell.fill = stripe_fills[(data_row - 2) % 2]
            if column.kind == "datetime" and value:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif column.kind == "date" and value:
                cell.number_format = "yyyy-mm-dd"
            elif column.kind == "number" and value != "":
                cell.number_format = "0.######"
            row.append(cell)
        sheet.append(row)
        row_count += 1

    sheet.auto_filter.ref = f"A1:{_column_letter(len(definition.columns))}{row_count}"
    return workbook


def preferred_column_width(column):
    """Keep identifiers readable without allowing long labels to create huge sheets."""
    label = column.label
    if column.kind == "bool":
        return 10
    if column.kind == "date":
        return 14
    if column.kind == "datetime":
        return 20
    if label in {"客户物料编码", "物料编码", "海关物料编码", "条码"}:
        return 24
    if any(token in label for token in ("客户代码", "供应商代码", "部门代码")):
        return 20
    if any(token in label for token in ("备注", "地址")):
        return 28
    if any(token in label for token in ("名称", "描述", "规格", "网址", "邮箱", "联系人", "职位", "电话", "传真")):
        return 24
    if any(token in label for token in ("科目", "帐号", "账号", "权限", "方式", "策略", "规则", "文件")):
        return 18
    if column.kind == "number":
        return 13
    return min(20, max(12, len(label) * 2 + 2))


def _column_letter(number):
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def workbook_response(definition, queryset):
    workbook = build_workbook(definition, queryset)
    output = BytesIO()
    workbook.save(output)
    filename = f"{definition.sheet_name}_{timezone.localtime():%Y%m%d_%H%M%S}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return response


def supplier_quote_workbook_response(queryset):
    workbook = build_workbook(SUPPLIER_QUOTE, queryset)
    sheet = workbook.create_sheet("报价阶梯明细")
    headers = (
        "报价单号", "行号", "数量大于等于", "材料单价", "加工单价", "含税单价", "未税单价",
        "备注", "录入人", "录入时间", "最后修改人", "最后修改时间",
    )
    widths = (18, 9, 16, 14, 14, 14, 14, 28, 14, 20, 14, 20)
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="1F2937")
    grid = Side(style="thin", color="C7D0DA")
    border = Border(left=grid, right=grid, top=grid, bottom=grid)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[_column_letter(index)].width = width
    header = []
    for value in headers:
        cell = WriteOnlyCell(sheet, value=value)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        header.append(cell)
    sheet.append(header)
    row_count = 1
    for supplier_quote in queryset.iterator(chunk_size=500):
        for tier in supplier_quote.lines.all():
            row_count += 1
            values = (
                supplier_quote.number, tier.line_number, tier.min_qty, tier.material_unit_price,
                tier.processing_unit_price, tier.unit_price, tier.untaxed_unit_price, tier.notes,
                tier.created_by, tier.created_at, tier.updated_by, tier.updated_at,
            )
            row = []
            for index, value in enumerate(values):
                cell = WriteOnlyCell(sheet, value=excel_value(value, "datetime" if index in {9, 11} else "number" if index in {1, 2, 3, 4, 5, 6} else "text"))
                cell.font = Font(name="Microsoft YaHei", size=10)
                cell.border = border
                if index in {9, 11} and value:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif index in {1, 2, 3, 4, 5, 6} and value is not None:
                    cell.number_format = "0.########"
                row.append(cell)
            sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{row_count}"
    output = BytesIO()
    workbook.save(output)
    filename = f"{SUPPLIER_QUOTE.sheet_name}_{timezone.localtime():%Y%m%d_%H%M%S}.xlsx"
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return response
