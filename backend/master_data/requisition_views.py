"""Three-level requisition workspace; legacy detail endpoints share the same guards."""
from decimal import Decimal
from io import BytesIO
from django.core.exceptions import ValidationError as ModelError
from django.db import transaction
from django.db.models import Q
from django.http import FileResponse
from django.utils import timezone
from openpyxl import Workbook
from rest_framework import serializers, decorators, response
from rest_framework.pagination import PageNumberPagination
from backend.models import (PurchaseRequisition, PurchaseRequisitionLine, PurchaseRequisitionSource, Material, Employee,
    AuditEvent, Partner, Department, PurchaseOrder, PurchaseOrderLine, SupplierQuote)
from backend.requisitions import assert_editable, material_snapshot
from .purchase_views import PurchaseRequisitionViewSet as LegacyRequisitionViewSet
from .api_common import EditableViewSet
from .exports import safe_text
from .serializers.purchase import PurchaseOrderLineSerializer, PurchaseOrderSerializer


class SourceSerializer(serializers.ModelSerializer):
    class Meta:
        model = PurchaseRequisitionSource
        fields = '__all__'


class LineSerializer(serializers.ModelSerializer):
    sources = SourceSerializer(many=True, read_only=True)
    class Meta:
        model = PurchaseRequisitionLine
        fields = '__all__'
        read_only_fields = ('converted_qty', 'calculation_snapshot', 'source_snapshot', 'material_snapshot',
            'source_kind', 'closed_by', 'closed_at', 'line_status', 'close_reason', 'source_object', 'source_detail_object',
            'created_by','created_at','updated_by','updated_at','modification_count')
        extra_kwargs = {'requisition': {'required': False}, 'mrp_demand_qty': {'required': False}}

    def validate(self, attrs):
        doc = attrs.get('requisition') or (self.instance.requisition if self.instance else self.context.get('document'))
        if doc:
            try:
                assert_editable(doc)
            except ModelError as exc:
                raise serializers.ValidationError(exc.messages)
        if self.instance:
            for field in ('requisition', 'line_number', 'material', 'uom', 'source_sales_order_line'):
                if field in attrs and attrs[field] != getattr(self.instance, field):
                    raise serializers.ValidationError('已保存行的身份及来源不能修改，请删除后重新增加')
        else:
            material = attrs.get('material')
            if not material or material.status != 'approved' or not material.active:
                raise serializers.ValidationError('请选择已审核且启用的物料')
            else:
                attrs['uom'] = material.uom
                attrs['material_snapshot'] = material_snapshot(material)
            attrs.setdefault('mrp_demand_qty', Decimal(0))
            if 'requested_qty' not in attrs and attrs['mrp_demand_qty']:
                attrs['source_kind'] = 'legacy_mrp'
                attrs['requested_qty'] = max(Decimal(0), attrs['mrp_demand_qty'] + attrs.get('safety_stock_qty', 0) - attrs.get('on_order_qty', 0) - attrs.get('available_qty', 0))
        quantity = attrs.get('requested_qty', getattr(self.instance, 'requested_qty', Decimal(0)))
        if quantity <= 0:
            raise serializers.ValidationError({'requested_qty': '请购数量必须大于0'})
        attrs.setdefault('needed_date', getattr(self.instance, 'needed_date', None) or (doc.needed_date if doc else None) or timezone.localdate())
        return attrs


class HeaderSerializer(serializers.ModelSerializer):
    lines = LineSerializer(many=True, read_only=True)

    def get_fields(self):
        fields = super().get_fields()
        if getattr(self.context.get('view'), 'action', '') == 'list':
            for key in ('lines', 'source_snapshot'):
                fields.pop(key, None)
        return fields
    class Meta:
        model = PurchaseRequisition
        fields = '__all__'
        read_only_fields = ('number', 'version', 'print_count', 'status', 'source_profile', 'source_snapshot',
            'created_by', 'created_at', 'updated_by', 'updated_at', 'approved_by', 'approved_at',
            'conversion_percent', 'is_confirmed', 'confirmed_by', 'confirmed_at', 'source_object')

    def validate(self, attrs):
        if attrs.get('requisition_type', getattr(self.instance, 'requisition_type', 'P')) != 'P':
            raise serializers.ValidationError('本期仅开放P采购请购维护')
        if attrs.get('document_type', getattr(self.instance, 'document_type', '')) != '':
            raise serializers.ValidationError('本期手工请购只开放普通单别')
        if self.instance and self.instance.source_profile:
            for key in ('site', 'requisition_type', 'document_type'):
                if key in attrs and attrs[key] != getattr(self.instance, key):
                    raise serializers.ValidationError('源单据站点及类型不可修改')
        user = self.context['request'].user
        employee_link = getattr(user, 'employee_link', None)
        attrs.setdefault('requester', getattr(self.instance, 'requester', '') or (employee_link.employee.code if employee_link else user.get_username()))
        if not attrs['requester'].strip():
            raise serializers.ValidationError('请购人不能为空')
        changed = not self.instance or attrs['requester'] != self.instance.requester
        if changed and attrs['requester'] != user.get_username() and not Employee.objects.filter(code=attrs['requester'], is_left=False, is_locked=False).exists():
            raise serializers.ValidationError('请购人不存在或已离职/锁定')
        return attrs


class RequisitionPagination(PageNumberPagination):
    page_size = 50


class PurchaseRequisitionViewSet(LegacyRequisitionViewSet):
    serializer_class = HeaderSerializer

    @decorators.action(detail=True, methods=['post'])
    @transaction.atomic
    def convert(self, request, pk=None):
        try:
            document = PurchaseRequisition.objects.select_for_update().get(pk=self.get_object().pk)
            if document.requisition_type != 'P':
                raise ModelError('本期仅支持P采购请购转采')
            return super().convert(request, pk)
        except ModelError as exc:
            raise serializers.ValidationError(exc.messages)

    def get_queryset(self):
        qs = PurchaseRequisition.objects.all().order_by('-request_date', '-number', '-id')
        if self.action != 'list':
            qs = qs.prefetch_related('lines__sources')
        q = self.request.query_params.get('q', '').strip()
        if q:
            qs = qs.filter(Q(number__icontains=q) | Q(requester__icontains=q) | Q(notes__icontains=q)
                | Q(lines__material_snapshot__pt_part__icontains=q) | Q(lines__source_snapshot__req_wo_nbr__icontains=q)).distinct()
        for field in ('status', 'requisition_type'):
            if self.request.query_params.get(field):
                qs = qs.filter(**{field:self.request.query_params[field]})
        from datetime import date
        for key, lookup in (('date_from','request_date__gte'),('date_to','request_date__lte')):
            if self.request.query_params.get(key):
                try:
                    qs = qs.filter(**{lookup:date.fromisoformat(self.request.query_params[key])})
                except ValueError:
                    raise serializers.ValidationError('日期格式错误')
        return qs

    def list(self, request, *args, **kwargs):
        if request.query_params.get('workspace') == '1':
            pager = RequisitionPagination()
            rows = pager.paginate_queryset(self.get_queryset(), request)
            fields = ('id','number','request_date','requester','requisition_type','document_type','status','version','source_profile')
            return pager.get_paginated_response([{key:getattr(row,key) for key in fields} for row in rows])
        return super().list(request, *args, **kwargs)

    @decorators.action(detail=False, methods=['get'])
    def navigate(self, request):
        ids = list(self.get_queryset().values_list('pk', flat=True))
        if not ids:
            raise serializers.ValidationError('当前查询范围没有单据')
        try:
            current = ids.index(int(request.query_params.get('current', 0)))
        except (ValueError, TypeError):
            current = 0
        direction = request.query_params.get('direction')
        target = {'first':0,'last':len(ids)-1,'previous':max(0,current-1),'next':min(len(ids)-1,current+1)}.get(direction)
        if target is None:
            raise serializers.ValidationError('导航方向无效')
        return response.Response(self.get_serializer(self.get_queryset().get(pk=ids[target])).data)

    @transaction.atomic
    def create(self, request, *args, **kwargs):
        return self.save_document(request)

    @transaction.atomic
    def update(self, request, *args, **kwargs):
        return self.save_document(request, self.get_object().pk, kwargs.get('partial', False))

    def save_document(self, request, pk=None, partial=False):
        doc = PurchaseRequisition.objects.select_for_update().get(pk=pk) if pk else None
        try:
            if doc:
                assert_editable(doc)
                if request.data.get('version') != doc.version:
                    raise serializers.ValidationError('单据已变更，请重新查询后保存')
            ser = self.get_serializer(doc, data=request.data, partial=partial)
            ser.is_valid(raise_exception=True)
            actor = request.user.get_username()
            doc = ser.save(**({'updated_by':actor} if pk else {'created_by':actor}))
            if 'lines' in request.data:
                entries = request.data['lines']
                if not isinstance(entries, list) or not entries:
                    raise serializers.ValidationError('请购单至少需要一行')
                existing = {r.pk:r for r in doc.lines.select_for_update()}
                kept, numbers = set(), set()
                next_number = max([0, *(r.line_number for r in existing.values())])
                for entry in entries:
                    values = dict(entry)
                    row_id = values.pop('id', None)
                    row = existing.get(row_id)
                    if row_id and not row:
                        raise serializers.ValidationError('明细行不属于当前单据')
                    if row_id in kept:
                        raise serializers.ValidationError('重复明细行')
                    if not row and not values.get('line_number'):
                        next_number += 10
                        values['line_number'] = next_number
                    number = values.get('line_number', row.line_number if row else None)
                    if not row and number in {r.line_number for r in existing.values()}:
                        raise serializers.ValidationError('已保存行号不能被新行复用，请使用更大的行号')
                    if number in numbers:
                        raise serializers.ValidationError('行号重复')
                    numbers.add(number)
                    values['requisition'] = doc.pk
                    line_ser = LineSerializer(row, data=values, partial=bool(row), context={'document':doc})
                    line_ser.is_valid(raise_exception=True)
                    saved = line_ser.save(updated_by=actor, updated_at=timezone.now(),
                        modification_count=row.modification_count+1 if row else 0,
                        **({'created_by':actor} if not row else {}))
                    kept.add(saved.pk)
                for row in existing.values():
                    if row.pk not in kept:
                        if row.rfq_lines.exists():
                            raise serializers.ValidationError('已被询价引用的行不能删除')
                        row.delete()
            doc.version += 1
            doc.save(update_fields=['version', 'updated_at'])
            AuditEvent.objects.create(model=doc._meta.label_lower, object_id=str(doc.pk), action='update' if pk else 'create', actor=actor,
                payload={'version':doc.version,'quantities':[{'line':r.line_number,'quantity':str(r.requested_qty)} for r in doc.lines.all()]})
            return response.Response(self.get_serializer(doc).data, status=200 if pk else 201)
        except ModelError as exc:
            raise serializers.ValidationError(exc.messages)

    @transaction.atomic
    def transition(self, obj, action, actor='', reason=''):
        doc = PurchaseRequisition.objects.select_for_update().get(pk=obj.pk)
        if doc.requisition_type != 'P':
            raise serializers.ValidationError('该历史类型只读')
        if action in ('void','restore','confirm','unconfirm','ratify','unratify','sales_confirm','sales_unconfirm'):
            raise serializers.ValidationError('请使用请购审核、反审核及行关闭操作')
        result = super().transition(doc, action, actor, reason)
        doc.refresh_from_db()
        doc.version += 1
        doc.save(update_fields=['version'])
        return response.Response(self.get_serializer(doc).data)

    @transaction.atomic
    def destroy(self, request, *args, **kwargs):
        doc = PurchaseRequisition.objects.select_for_update().get(pk=self.get_object().pk)
        try:
            assert_editable(doc)
            if doc.rfqs.exists():
                raise ModelError('已被询价引用的请购不能删除')
            AuditEvent.objects.create(model=doc._meta.label_lower, object_id=str(doc.pk), action='delete', actor=request.user.get_username(),
                payload={'number':doc.number,'source_profile':doc.source_profile})
            doc.delete()
        except ModelError as exc:
            raise serializers.ValidationError(exc.messages)
        return response.Response(status=204)

    @decorators.action(detail=False, methods=['get'])
    def lookups(self, request):
        q = request.query_params.get('q','').strip()
        if request.query_params.get('kind') == 'employee':
            rows = Employee.objects.filter(is_left=False,is_locked=False).filter(Q(code__icontains=q)|Q(name__icontains=q))[:50]
            return response.Response([{'id':r.code,'code':r.code,'name':r.name} for r in rows])
        if request.query_params.get('kind') == 'supplier':
            rows = Partner.objects.filter(kind__in=['supplier', 'both'], status='approved').filter(
                Q(code__icontains=q) | Q(name__icontains=q) | Q(short_name__icontains=q)
            ).select_related('currency')[:50]
            return response.Response([{'id':r.pk,'code':r.code,'name':r.short_name or r.name,
                'currency':r.currency_id,'currency_code':r.currency.code if r.currency_id else '',
                'payment_method':r.payment_method,'tax_rate':str(r.tax_rate)} for r in rows])
        rows = Material.objects.filter(status='approved',active=True).filter(Q(code__icontains=q)|Q(name__icontains=q)).select_related('uom','category')[:50]
        result = [{'id':r.pk,'material':r.pk,'code':r.code,'name':r.name,'uom':r.uom_id,'snapshot':material_snapshot(r)} for r in rows]
        return response.Response(result)

    @decorators.action(detail=True, methods=['post'], url_path='convert-purchase')
    @transaction.atomic
    def convert_purchase(self, request, pk=None):
        """Create ordinary purchase orders directly from remaining P requisition quantities."""
        requisition = PurchaseRequisition.objects.select_for_update().get(pk=pk)
        if requisition.requisition_type != 'P' or requisition.status != 'approved':
            raise serializers.ValidationError('只有已审核P采购请购单可以转采购单')
        supplier_id = request.data.get('supplier_id')
        supplier = Partner.objects.select_for_update().select_related('currency').filter(
            pk=supplier_id, kind__in=['supplier', 'both'], status='approved').first()
        if not supplier:
            raise serializers.ValidationError({'supplier_id': '请选择已审核供应商'})
        department_id = request.data.get('department_id') or requisition.department_id
        department = Department.objects.filter(pk=department_id).first()
        if not department:
            raise serializers.ValidationError({'department_id': '请指定采购部门'})
        if not supplier.currency_id:
            raise serializers.ValidationError({'supplier_id': '供应商未设置常用币种'})
        requested = request.data.get('lines') or []
        if not isinstance(requested, list) or not requested:
            raise serializers.ValidationError({'lines': '至少选择一条可转采购明细'})
        line_ids = [item.get('line_id') for item in requested]
        if len(line_ids) != len(set(line_ids)) or any(not value for value in line_ids):
            raise serializers.ValidationError({'lines': '请购明细不能重复选择'})
        rows = list(PurchaseRequisitionLine.objects.select_for_update().select_related('material','uom').filter(
            requisition=requisition, pk__in=line_ids).order_by('line_number'))
        if len(rows) != len(line_ids):
            raise serializers.ValidationError({'lines': '存在不属于当前请购单的明细'})
        by_id = {row.pk: row for row in rows}
        values = []
        for item in requested:
            row = by_id[item['line_id']]
            remaining = row.requested_qty - row.converted_qty
            quote = SupplierQuote.objects.filter(supplier=supplier, material=row.material,
                quote_type='purchase', status='approved', is_confirmed=True, is_ratified=True,
                effective_date__lte=timezone.localdate()).filter(
                    Q(expiry_date__isnull=True) | Q(expiry_date__gte=timezone.localdate())
                ).order_by('-effective_date', '-updated_at').first()
            try:
                quantity = Decimal(str(item.get('quantity', remaining)))
                raw_price = item.get('unit_price')
                price = Decimal(str(raw_price)) if raw_price not in (None, '') else (quote.price_for(quantity)[0] if quote else Decimal('0'))
            except Exception as exc:
                raise serializers.ValidationError({'lines': '数量和单价必须是数字'}) from exc
            if row.line_status == 'C' or not row.material_id:
                raise serializers.ValidationError({'lines': f'请购行{row.line_number}已关闭或物料未映射本地主档'})
            if quantity <= 0 or quantity > remaining:
                raise serializers.ValidationError({'lines': f'请购行{row.line_number}可转数量为{remaining}'})
            if price < 0:
                raise serializers.ValidationError({'lines': f'请购行{row.line_number}单价不能为负数'})
            promised = item.get('promised_date') or row.needed_date or requisition.request_date
            values.append((row, quantity, price, promised, quote))
        order = PurchaseOrder.objects.create(
            supplier=supplier, currency=supplier.currency, department=department,
            order_date=timezone.localdate(), buyer=request.user.get_username(),
            payment_method=supplier.payment_method or '', tax_rate=supplier.tax_rate or Decimal('0'),
            promised_date=max(value[3] for value in values), purchase_type=PurchaseOrder.PurchaseType.PRODUCTION,
            source_requisition=requisition, created_by=request.user.get_username(),
        )
        for index, (row, quantity, price, promised, quote) in enumerate(values, 1):
            line = PurchaseOrderLine(order=order, line_number=index * 10, material=row.material, uom=row.uom,
                quantity=quantity, unit_price=price, promised_date=promised, source_requisition_line=row,
                source_supplier_quote=quote)
            line.full_clean()
            line.save()
        requisition._prefetched_objects_cache = {}
        total = sum((line.requested_qty for line in requisition.lines.all()), Decimal('0'))
        applied = sum((line.converted_qty for line in requisition.lines.all()), Decimal('0'))
        requisition.conversion_percent = min(Decimal('100'), applied * Decimal('100') / total) if total else Decimal('0')
        requisition.version += 1
        requisition.save(update_fields=['conversion_percent', 'version', 'updated_at'])
        AuditEvent.objects.create(model=order._meta.label_lower, object_id=str(order.pk), action='convert_from_requisition',
            actor=request.user.get_username() or 'system', payload={'source_requisition_id': requisition.pk,
            'source_line_ids': line_ids, 'supplier_id': supplier.pk})
        payload = PurchaseOrderSerializer(order).data
        payload['lines'] = [dict(PurchaseOrderLineSerializer(line).data, source_requisition_number=requisition.number,
            source_requisition_line=line.source_requisition_line.line_number) for line in order.lines.select_related('source_requisition_line')]
        return response.Response(payload, status=201)

    @decorators.action(detail=True, methods=['post'])
    @transaction.atomic
    def close_line(self, request, pk=None):
        doc = PurchaseRequisition.objects.select_for_update().get(pk=self.get_object().pk)
        if doc.requisition_type != 'P' or doc.status != 'approved':
            raise serializers.ValidationError('仅已审核P采购请购可以关闭行')
        reason = str(request.data.get('reason','')).strip()
        if not reason:
            raise serializers.ValidationError('请输入关闭原因')
        line = doc.lines.filter(pk=request.data.get('line_id')).first()
        if not line:
            raise serializers.ValidationError('请选择当前单据行')
        line.line_status, line.closed_by, line.closed_at, line.close_reason = 'C', request.user.get_username(), timezone.now(), reason
        line.save(update_fields=['line_status','closed_by','closed_at','close_reason'])
        doc.version += 1
        doc.save(update_fields=['version'])
        AuditEvent.objects.create(model=line._meta.label_lower,object_id=str(line.pk),action='close',actor=request.user.get_username(),reason=reason)
        return response.Response(self.get_serializer(doc).data)

    @decorators.action(detail=True, methods=['post'])
    @transaction.atomic
    def print(self, request, pk=None):
        doc = PurchaseRequisition.objects.select_for_update().get(pk=self.get_object().pk)
        doc.print_count += 1
        doc.save(update_fields=['print_count'])
        AuditEvent.objects.create(model=doc._meta.label_lower, object_id=str(doc.pk),action='print',actor=request.user.get_username())
        return response.Response(self.get_serializer(doc).data)

    @decorators.action(detail=True, methods=['post'])
    def export(self, request, pk=None):
        doc = self.get_object()
        book = Workbook()
        sheet = book.active
        sheet.title = '表头'
        for label, value in [('请购单号',doc.number),('日期',str(doc.request_date)),('类型',doc.requisition_type),('单别',doc.document_type),('请购人',doc.requester),('备注',doc.notes),('审核人',doc.approved_by)]:
            sheet.append([label,safe_text(value)])
        sheet = book.create_sheet('物料明细')
        sheet.append(['行号','物料','名称','规格','单位','需求日期','请购数量','已采购数量','状态','备注'])
        source_sheet = book.create_sheet('合并请购明细')
        source_fields = [('reqd_wod_qty_req_std','标准需求数量（领料单位）'),('reqd_wod_qty_req_scr','损耗需求数量（领料单位）'),
            ('reqd_wod_qty_req','需求数量（领料单位）'),('reqd_mrp_req_tot','MRP总需求'),('reqd_mrp_ld_net','MRP可用库存'),
            ('reqd_mrp_ord','MRP订购中/在制'),('reqd_mrp_pr','MRP请购中'),('reqd_mrp_safe','MRP安全库存'),
            ('reqd_mrp_alloc','MRP分配数量'),('reqd_mrp_left','MRP库存余量'),('reqd_mrp_short','MRP短缺数量'),
            ('reqd_wod_part','主件编码'),('parent_name','主件名称'),('reqd_wo_qty_ord','生产数量')]
        source_sheet.append(['请购行','序号','来源单号','标志','来源序号','建议请购数',*[label for _,label in source_fields],'源字段快照'])
        raw_sheet = book.create_sheet('物料及MRP字段')
        raw_rows = list(doc.lines.prefetch_related('sources'))
        raw_keys = sorted({k for r in raw_rows for k in (r.material_snapshot.keys() | r.source_snapshot.keys())})
        raw_sheet.append(['行号',*raw_keys])
        import json
        for line in raw_rows:
            p = line.material_snapshot
            combined = {**p,**line.source_snapshot}
            raw_sheet.append([line.line_number,*[safe_text(str(combined.get(k,''))) for k in raw_keys]])
            sheet.append([line.line_number,*[safe_text(p.get(k,'')) for k in ('pt_part','pt_desc1','pt_spec','pt_um')],str(line.needed_date or ''),line.requested_qty,line.converted_qty,line.line_status,safe_text(line.notes)])
            for source in line.sources.all():
                source_sheet.append([line.line_number,source.sequence,safe_text(source.source_number),safe_text(source.source_marker),source.source_sequence,source.quantity,
                    *[safe_text(str(source.snapshot.get(k,''))) for k,_ in source_fields],safe_text(json.dumps(source.snapshot,ensure_ascii=False))])
        output = BytesIO()
        book.save(output)
        output.seek(0)
        return FileResponse(output, as_attachment=True, filename=doc.number+'.xlsx')


class PurchaseRequisitionLineViewSet(EditableViewSet):
    queryset = PurchaseRequisitionLine.objects.select_related('requisition','material').prefetch_related('sources')
    serializer_class = LineSerializer

    @transaction.atomic
    def perform_create(self, serializer):
        doc = PurchaseRequisition.objects.select_for_update().get(pk=serializer.validated_data['requisition'].pk)
        try:
            assert_editable(doc)
        except ModelError as exc:
            raise serializers.ValidationError(exc.messages)
        serializer.save(created_by=self.request.user.get_username(), updated_by=self.request.user.get_username())
        doc.version += 1
        doc.save(update_fields=['version'])

    @transaction.atomic
    def perform_update(self, serializer):
        doc = PurchaseRequisition.objects.select_for_update().get(pk=serializer.instance.requisition_id)
        try:
            assert_editable(doc)
        except ModelError as exc:
            raise serializers.ValidationError(exc.messages)
        serializer.save(updated_by=self.request.user.get_username(),updated_at=timezone.now(),
            modification_count=serializer.instance.modification_count+1)
        doc.version += 1
        doc.save(update_fields=['version'])

    @transaction.atomic
    def destroy(self, request, *args, **kwargs):
        line = self.get_object()
        doc = PurchaseRequisition.objects.select_for_update().get(pk=line.requisition_id)
        try:
            assert_editable(doc)
            if line.rfq_lines.exists():
                raise ModelError('已被询价引用，不能删除')
            line.delete()
            doc.version += 1
            doc.save(update_fields=['version'])
        except ModelError as exc:
            raise serializers.ValidationError(exc.messages)
        return response.Response(status=204)


class PurchaseOrderLineViewSet(EditableViewSet):
    queryset = PurchaseOrderLine.objects.all()
    serializer_class = PurchaseOrderLineSerializer

    @transaction.atomic
    def perform_create(self, serializer):
        from .api_common import ensure_editable
        ensure_editable(serializer.validated_data['order'])
        try:
            serializer.save()
        except ModelError as exc:
            raise serializers.ValidationError(exc.messages)

    @transaction.atomic
    def perform_update(self, serializer):
        try:
            return super().perform_update(serializer)
        except ModelError as exc:
            raise serializers.ValidationError(exc.messages)
