import json
from decimal import Decimal, InvalidOperation
from tempfile import SpooledTemporaryFile

from django.http import FileResponse
from openpyxl import Workbook
from rest_framework import decorators, response, serializers, viewsets
from backend.domain.inventory_analysis import InventoryAnalysisSnapshot
from backend.inventory_analysis import (BULK_COLUMNS, BULK_NUMBERS, MONTHLY_COLUMNS, MONTHLY_NUMBERS,
    AVERAGES, DERIVED, GROUPS, aggregate_monthly, decimal)
from .api_common import ErpRolePermission
from .inventory_snapshot_views import integer, safe_text


from functools import lru_cache
from django.db import connection


@lru_cache(maxsize=4)
def snapshot_payload(database, pk, version):
    # Cached objects are immutable to callers; bulk_payload copies its selected rows.
    return InventoryAnalysisSnapshot.objects.get(pk=pk).payload


def bulk_source(site, source, part=None):
    if part is not None:
        from backend.inventory_bulk_paging import select_bulk, live_rows
        snapshot, batch = select_bulk(site, source)
        return (snapshot, live_rows(snapshot, batch, [part])) if snapshot else (None, {})
    from backend.live_inventory import live_batch, bulk_payload
    batch = live_batch() if source != 'dgyzx1' else None
    qs = InventoryAnalysisSnapshot.objects.defer('payload')
    snapshot = qs.filter(kind='localbulk', site=site).first() if batch else None
    snapshot = snapshot or qs.filter(kind='bulk', site=site).first()
    if snapshot is None:
        return None, {}
    snapshot.payload = snapshot_payload(str(connection.settings_dict['NAME']), snapshot.pk, snapshot.captured_at)
    if part is not None:
        part = next((r['tmp_part'] for r in snapshot.payload.get('rows', []) if r['tmp_part'].casefold() == part.casefold()), part)
    snapshot.report_source = 'local' if batch else 'coerp_dgyzx1'
    return snapshot, bulk_payload(snapshot, batch, part=part) if batch else snapshot.payload


def numeric(value):
    try:
        result = Decimal(str(value))
        if not result.is_finite():
            raise ValueError()
        return result
    except (InvalidOperation, ValueError, TypeError):
        raise serializers.ValidationError('数量条件必须为有限数值')


class InventoryAnalysisViewSet(viewsets.ViewSet):
    permission_classes = (ErpRolePermission,)
    kind = 'bulk'

    def query(self, params):
        monthly = self.kind == 'monthly'
        columns, numbers = (MONTHLY_COLUMNS, MONTHLY_NUMBERS) if monthly else (BULK_COLUMNS, BULK_NUMBERS)
        period = params.get('period', '') if monthly else ''
        if monthly:
            snapshot = InventoryAnalysisSnapshot.objects.filter(kind=self.kind, site=params.get('site', '1000'), period=period).first()
            payload = snapshot.payload if snapshot else {}
        else:
            low, high = params.get('material_from', '').strip(), params.get('material_to', '').strip()
            part = low if low and low == high else None
            snapshot, payload = bulk_source(params.get('site', '1000'), params.get('source', 'local'), part=part)
        if snapshot is None:
            return [], columns, numbers, None
        # ponytail: aggregate the ~37k-row capture in memory; use indexed numeric rows if capture size outgrows request latency.
        rows = payload['rows']
        ranges = [('material', 'inb_part' if monthly else 'tmp_part')]
        if monthly:
            ranges += [('category', 'pt_prod_line'), ('group', 'pt_group'), ('part_type', 'pt_part_type'), ('location', 'inb_loc')]
        for name, field in ranges:
            low, high = params.get(name+'_from', '').strip().casefold(), params.get(name+'_to', '').strip().casefold()
            if low and high and low > high:
                raise serializers.ValidationError(f'{name}: 起始编码不能大于结束编码')
            rows = [r for r in rows if (not low or str(r.get(field, '')).casefold() >= low) and (not high or str(r.get(field, '')).casefold() <= high)]
        query = params.get('q', '').strip().casefold()
        if query:
            def description(row):
                return payload['masters']['materials'].get(row['inb_part'], {}) if monthly else row
            rows = [r for r in rows if any(query in str(description(r).get(k, '')).casefold() for k in ('pt_desc1', 'pt_spec'))]
        if monthly:
            for name, field in [('locations', 'inb_loc'), ('location_types', 'loc_type_code'), ('cost_types', 'loc_cost_type')]:
                chosen = {v.strip() for v in params.get(name, '').split(',') if v.strip()}
                if chosen:
                    rows = [r for r in rows if r.get(field) in chosen]
            groups = [g for g in params.get('groups', 'inb_part,loc_cost_type').split(',') if g]
            if any(g not in GROUPS for g in groups):
                raise serializers.ValidationError('未知汇总维度')
            rows = aggregate_monthly(rows, groups, payload['masters'])
            threshold = numeric(params.get('diff_rate') or 0)
            if threshold:
                rows = [r for r in rows if decimal(r['diff_rate']) >= threshold]
        try:
            conditions = json.loads(params.get('numeric_filters', '{}'))
        except (ValueError, TypeError):
            raise serializers.ValidationError('数量检索格式错误')
        if not isinstance(conditions, dict):
            raise serializers.ValidationError('数量检索格式错误')
        for field, condition in conditions.items():
            if field not in numbers or not isinstance(condition, dict) or condition.get('op') not in {'eq', 'ne', 'gt', 'ge', 'lt', 'le', 'between'}:
                raise serializers.ValidationError('未知数量字段或比较方式')
            value, op = numeric(condition.get('value')), condition['op']
            high = numeric(condition.get('high')) if op == 'between' else value
            if op == 'between' and value > high:
                raise serializers.ValidationError('数量下限不能大于上限')
            def matches(row):
                actual = decimal(row.get(field))
                return {'eq': actual == value, 'ne': actual != value, 'gt': actual > value, 'ge': actual >= value,
                        'lt': actual < value, 'le': actual <= value, 'between': value <= actual <= high}[op]
            rows = [r for r in rows if matches(r)]
        sort = params.get('sort') or ('inb_part' if monthly else 'tmp_part')
        if sort not in dict(columns) or params.get('direction', 'asc') not in {'asc', 'desc'}:
            raise serializers.ValidationError('未知排序字段或方向')
        rows = sorted(rows, key=lambda r: decimal(r.get(sort)) if sort in numbers else str(r.get(sort, '')).casefold(), reverse=params.get('direction') == 'desc')
        return rows, columns, numbers, snapshot

    def list(self, request):
        params = request.query_params
        if self.kind == 'bulk' and params.get('page_mode') == '1' and params.get('sort', 'tmp_part') == 'tmp_part' and params.get('numeric_filters', '{}') in ('', '{}'):
            return self.bulk_page(params)
        rows, columns, numbers, snapshot = self.query(request.query_params)
        page, size = integer(request.query_params, 'page', 1), integer(request.query_params, 'page_size', 100, 500)
        totals = {k: str(sum((decimal(r.get(k)) for r in rows), Decimal(0))) for k in numbers - (AVERAGES | DERIVED if self.kind == 'monthly' else set())}
        return response.Response({'columns': columns, 'numeric_columns': sorted(numbers), 'count': len(rows),
            'results': rows[(page-1)*size:page*size], 'totals': totals, 'page': page,
            'snapshot': {'id': snapshot.id, 'captured_at': snapshot.captured_at,
                         'source': getattr(snapshot, 'report_source', 'coerp_dgyzx1'),
                         'consistent_snapshot': snapshot.payload.get('consistent_snapshot')} if snapshot else None})

    def bulk_page(self, params):
        from backend.inventory_bulk_paging import select_bulk, page_rows
        page, size = integer(params, 'page', 1), integer(params, 'page_size', 100, 500)
        low, high = (params.get('material_' + suffix, '').strip().casefold() for suffix in ('from', 'to'))
        if low and high and low > high:
            raise serializers.ValidationError('material: 起始编码不能大于结束编码')
        if params.get('direction', 'asc') not in ('asc', 'desc'):
            raise serializers.ValidationError('未知排序方向')
        snapshot, batch = select_bulk(params.get('site', '1000'), params.get('source', 'local'))
        rows, count = page_rows(snapshot, batch, params, page, size) if snapshot else ([], 0)
        return response.Response({'columns': BULK_COLUMNS, 'numeric_columns': sorted(BULK_NUMBERS),
            'results': rows, 'count': count, 'page': page, 'page_size': size, 'totals_scope': 'page',
            'totals': {k: str(sum((decimal(r.get(k)) for r in rows), Decimal(0))) for k in BULK_NUMBERS},
            'snapshot': {'id': snapshot.pk, 'captured_at': snapshot.captured_at,
                'source': snapshot.report_source} if snapshot else None})

    @decorators.action(detail=False, methods=['get'])
    def choices(self, request):
        if self.kind == 'bulk':
            return response.Response({'periods': [], 'cost_types': {}, 'location_types': {}})
        site = request.query_params.get('site', '1000')
        snapshots = InventoryAnalysisSnapshot.objects.filter(kind=self.kind, site=site).order_by('-period')
        latest = snapshots.first()
        masters = latest.payload.get('masters', {}) if latest else {}
        return response.Response({'periods': list(snapshots.values_list('period', flat=True)),
            'cost_types': masters.get('cost_types', {}),
            'location_types': {r['loc_type']: r['loc_type_name'] for r in masters.get('locations', {}).values()}})

    @decorators.action(detail=False, methods=['get'])
    def details(self, request):
        snapshot, payload = bulk_source(request.query_params.get('site', '1000'), request.query_params.get('source', 'local'), part=request.query_params.get('part', ''))
        tab, part = request.query_params.get('tab', 'stock'), request.query_params.get('part', '')
        if tab not in {'stock', 'demand', 'order', 'requisition', 'allocation'}:
            raise serializers.ValidationError('未知明细类型')
        rows = [r for r in payload.get('details', []) if r['part'].casefold() == part.casefold() and r['tab'] == tab
                and (tab != 'stock' or request.query_params.get('include_wip') == '1' or r['location_type'] != 'WIP')] if snapshot else []
        return response.Response({'results': rows, 'total': str(sum((decimal(r['quantity']) for r in rows), Decimal(0)))})

    @decorators.action(detail=False, methods=['post'])
    def export(self, request):
        params = request.data.get('filters', {})
        if not isinstance(params, dict):
            raise serializers.ValidationError('查询条件格式错误')
        rows, columns, numbers, snapshot = self.query(params)
        if snapshot is None:
            raise serializers.ValidationError('该期间尚未同步')
        book = Workbook(write_only=True)
        sheet = book.create_sheet('查询结果')
        sheet.append([label for _, label in columns])
        for row in rows:
            sheet.append([decimal(row.get(k)) if k in numbers else safe_text(row.get(k, '')) for k, _ in columns])
        info = book.create_sheet('查询条件')
        info.append(['数据源', getattr(snapshot, 'report_source', 'coerp_dgyzx1')])
        info.append(['同步时间', str(snapshot.captured_at)])
        for key, value in params.items():
            info.append([safe_text(key), safe_text(str(value))])
        output = SpooledTemporaryFile(max_size=8*1024*1024)
        book.save(output)
        output.seek(0)
        return FileResponse(output, as_attachment=True, filename='inventory-'+self.kind+'.xlsx')


class MaterialInventoryBatchViewSet(InventoryAnalysisViewSet):
    kind = 'bulk'


class MaterialMonthlyInventoryViewSet(InventoryAnalysisViewSet):
    kind = 'monthly'
