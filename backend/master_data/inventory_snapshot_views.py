"""Indexed read-only source reports and bounded lookup endpoints."""
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from tempfile import SpooledTemporaryFile

from django.db import connection
from django.db.models import DecimalField, Q, Subquery
from django.db.models.functions import Cast
from django.http import FileResponse
from django.utils import timezone
from openpyxl import Workbook
from rest_framework import decorators, response, serializers

from backend.domain.inventory_reports import InventoryReportBatch, InventoryReportLookup, InventoryReportRow
from .exports import safe_text

INVENTORY_COLUMNS = [
    ("location_code", "库位"), ("location_name", "库位名称"), ("location_type_name", "库位类型"),
    ("material_code", "物料编码"), ("material_name", "物料名称"), ("specification", "规格"),
    ("quantity", "库存数量"), ("uom_name", "单位"), ("batch_number", "批号"),
    ("storage_position", "默认储位"), ("standard_cost", "单位标准成本"), ("inventory_amount", "库存金额"),
    ("category_code", "产品类"), ("category_name", "产品类名称"), ("group_code", "产品组"),
    ("group_name", "产品组名称"), ("part_type", "零件类型"), ("part_type_name", "零件类型名称"),
    ("sales_price", "单位不含税售价"), ("sales_amount", "不含税销售金额"),
    ("customer_material_code", "客户物料编码"), ("currency", "售价币种"), ("price_status", "售价状态"),
]
MOVEMENT_COLUMNS = [
    ("material_code", "物料编码"), ("material_name", "物料名称"), ("specification", "物料规格"),
    ("category_code", "产品类"), ("category_name", "产品类名称"), ("group_code", "产品组"),
    ("group_name", "产品组名称"), ("part_type", "零件类型"), ("part_type_name", "零件类型名称"),
    ("keeper_code", "仓管员"), ("keeper_name", "仓管员姓名"), ("movement_type", "进出类型"),
    ("movement_name", "进出类型描述"), ("document_date", "单据日期"), ("posted_at", "审核时间"),
    ("location_code", "库位"), ("location_name", "库位名称"), ("location_type_name", "库位类型"),
    ("batch_number", "批号"), ("receipt_quantity", "入库数量"), ("issue_quantity", "出库数量"),
    ("uom_name", "库存单位"), ("order_number", "订单号"), ("order_line", "订单行号"),
    ("flag", "标志"), ("display_document_number", "单据编号"), ("document_line", "单据行号"),
    ("partner_code", "往来对象编号"), ("partner_name", "供应商/客户/部门"),
    ("unit_cost", "单位成本"), ("movement_amount", "进出金额"), ("transaction_price", "订单单价"),
    ("currency", "订单币种"), ("notes", "备注"),
]
NUMERIC_COLUMNS = {"quantity", "receipt_quantity", "issue_quantity", "standard_cost", "inventory_amount",
                   "sales_price", "sales_amount", "unit_cost", "movement_amount", "transaction_price"}
RANGES = {"material": "material_code", "category": "category_code", "location": "location_code",
          "lot": "batch_number", "group": "group_code", "part_type": "part_type"}


def integer(params, key, default, maximum=2147483647):
    try:
        value = int(params.get(key, default))
        if value < 1 or (maximum and value > maximum):
            raise ValueError
        return value
    except (ValueError, TypeError):
        raise serializers.ValidationError({key: "请输入有效正整数"})


def flag(params, key, default):
    value = str(params.get(key, "1" if default else "0")).lower()
    if value not in {"0", "1", "true", "false"}:
        raise serializers.ValidationError({key: "布尔条件格式错误"})
    return value in {"1", "true"}


def batch_for(params):
    batches = InventoryReportBatch.objects.all()
    if params.get("batch_id"):
        batch = batches.filter(pk=integer(params, "batch_id", 1)).first()
        if batch is None:
            raise serializers.ValidationError("查询批次不存在，请重新查询")
        if params.get('source') == 'local' and batch.source != 'local':
            return batches.filter(source='local', active=True).first() or batch
        if params.get('source') == 'dgyzx1' and batch.source == 'local':
            raise serializers.ValidationError('所选批次为本地业务账，请重新选择历史查询')
        return batch
    if params.get('source') == 'dgyzx1':
        return batches.exclude(source='local').order_by('-id').first()
    return batches.filter(active=True).order_by("-id").first()


def range_filter(qs, batch, params, kind, field):
    lower, upper = (str(params.get(f"{kind}_{suffix}", "")).strip() for suffix in ("from", "to"))
    if not lower and not upper:
        return qs
    if batch.source == 'local':
        if lower and upper and lower.upper() > upper.upper():
            raise serializers.ValidationError({kind: '起始编码不能大于结束编码'})
        if lower:
            qs = qs.filter(**{field + '__gte': lower.upper() if lower.isascii() else lower})
        if upper:
            qs = qs.filter(**{field + '__lte': upper.upper() if upper.isascii() else upper})
        return qs
    codes = InventoryReportLookup.objects.filter(batch=batch, kind=kind)
    if kind in {"location", "lot"} and params.get("site"):
        codes = codes.filter(site=params["site"])
    bounds = []
    matched = []
    # Source-order ranks retain SQL Server collation for Chinese dictionary codes.
    for value, end in ((lower, False), (upper, True)):
        if not value:
            bounds.append(None)
            matched.append(False)
            continue
        match = codes.filter(code__iexact=value).order_by("source_order").first()
        matched.append(match is not None)
        if match:
            rank = match.source_order
        elif value.isascii():
            candidates = codes.filter(**{"code__lte" if end else "code__gte": value.upper()})
            candidate = candidates.order_by("-source_order" if end else "source_order").first()
            rank = candidate.source_order if candidate else (-1 if end else 2147483647)
        else:
            raise serializers.ValidationError({kind: "中文范围边界请从检索结果选择，以保持原库排序"})
        bounds.append(rank)
    if lower and upper and ((lower.isascii() and upper.isascii() and lower.upper() > upper.upper()) or
                            (all(matched) and bounds[0] > bounds[1])):
        raise serializers.ValidationError({kind: "起始编码不能大于结束编码"})
    if bounds[0] is not None:
        codes = codes.filter(source_order__gte=bounds[0])
    if bounds[1] is not None:
        codes = codes.filter(source_order__lte=bounds[1])
    return qs.filter(**{field + "__in": Subquery(codes.values("code"))})


def filtered_rows(batch, kind, params):
    qs = InventoryReportRow.objects.filter(batch=batch, kind=kind)
    if params.get("site"):
        qs = qs.filter(site=params["site"])
    for key, field in RANGES.items():
        if kind == "inventory" and key not in {"material", "category", "location", "lot"}:
            continue
        qs = range_filter(qs, batch, params, key, field)
    text = str(params.get("description", "")).strip()
    if text:
        qs = qs.filter(Q(material_name__icontains=text) | Q(specification__icontains=text))
    if kind == "inventory":
        if flag(params, "nonzero_quantity", False):
            qs = qs.annotate(stock_quantity=Cast('payload__quantity', DecimalField(max_digits=28, decimal_places=8))).exclude(stock_quantity=0)
        if not flag(params, "include_wip", False):
            qs = qs.exclude(location_type="WIP")
        return qs.order_by("location_code", "material_code", "batch_number", "id")
    for param, field in {"location": "location_code", "lot": "batch_number", "keeper": "keeper_code",
                         "movement_type": "movement_type", "document_number": "document_number",
                         "order_number": "order_number", "partner": "partner_code", "partner_kind": "partner_kind"}.items():
        if params.get(param):
            qs = qs.filter(**{field + "__iexact": str(params[param]).strip()})
    for prefix, field in (("document", "document_date"), ("posted", "posted_at")):
        parsed = []
        for suffix in ("from", "to"):
            raw = params.get(f"{prefix}_{suffix}")
            try:
                parsed.append(date.fromisoformat(raw) if raw else None)
            except (ValueError, TypeError):
                raise serializers.ValidationError({prefix: "日期格式错误"})
        start, end = parsed
        if start and end and start > end:
            raise serializers.ValidationError({prefix: "开始日期不能晚于结束日期"})
        if field == "posted_at":
            if start:
                qs = qs.filter(posted_at__gte=timezone.make_aware(datetime.combine(start, time.min)))
            if end:
                if end == date.max:
                    raise serializers.ValidationError({prefix: "结束日期超出支持范围"})
                qs = qs.filter(posted_at__lt=timezone.make_aware(datetime.combine(end + timedelta(days=1), time.min)))
        else:
            if start:
                qs = qs.filter(document_date__gte=start)
            if end:
                qs = qs.filter(document_date__lte=end)
    for parameter, types in (("exclude_cost", ("501",)), ("exclude_line", ("170", "270"))):
        if flag(params, parameter, True):
            if params.get("movement_type") in types:
                raise serializers.ValidationError({parameter: "所选进出类型被排除，请取消对应排除条件"})
            qs = qs.exclude(movement_type__in=types)
    sort = params.get("sort", "document_date")
    if sort not in {"document_date", "posted_at", "order_number"}:
        raise serializers.ValidationError({"sort": "未知排序方式"})
    return qs.order_by("material_code", sort, "source_table", "source_key", "id")


class SnapshotReportMixin:
    report_kind = "inventory"

    @decorators.action(detail=False, methods=["get"])
    def availability(self, request):
        if InventoryReportBatch._meta.db_table not in connection.introspection.table_names():
            return response.Response({"ready": False})
        return response.Response({"ready": InventoryReportBatch.objects.filter(active=True).exists()})

    def snapshot_list(self, request):
        params = request.query_params
        batch = batch_for(params)
        page = integer(params, "page", 1)
        size = integer(params, "page_size", 100, 500)
        columns = INVENTORY_COLUMNS if self.report_kind == "inventory" else MOVEMENT_COLUMNS
        if batch is None:
            return response.Response({"count": 0, "results": [], "columns": columns, "batch": None, "page": page})
        qs = filtered_rows(batch, self.report_kind, params)
        total = qs.count()
        records = [{**row.payload, "id": row.id} for row in qs[(page - 1) * size:page * size]]
        metadata = {key: batch.manifest.get(key) for key in
                    ("started_at", "finished_at", "date_from", "date_to", "isolation", "consistent_snapshot")}
        return response.Response({"count": total, "results": records, "columns": columns,
                                  "page": page, "page_size": size,
                                  "batch": {"id": batch.id, "source": batch.source,
                                            "created_at": batch.created_at, **metadata}})

    @decorators.action(detail=False, methods=["get"])
    def lookup(self, request):
        params = request.query_params
        batch = batch_for(params)
        kind = params.get("kind", "material")
        if kind not in {*RANGES, "keeper", "movement_type", "customer", "supplier", "department"}:
            raise serializers.ValidationError("未知检索类型")
        page = integer(params, "page", 1)
        size = 50
        qs = InventoryReportLookup.objects.filter(batch=batch, kind=kind)
        if kind in {"location", "lot"} and params.get("site"):
            qs = qs.filter(site=params["site"])
        if kind == "lot" and batch:
            row_params = {k: params[k] for k in ("site", "material_from", "material_to", "location_from", "location_to", "location") if params.get(k)}
            row_params.update(exclude_cost="0", exclude_line="0", include_wip="1")
            qs = qs.filter(code__in=Subquery(filtered_rows(batch, self.report_kind, row_params).order_by().values("batch_number")))
        field = params.get("field", "code")
        match = params.get("match", "contains")
        if field not in {"code", "name", "specification"} or match not in {"contains", "equals"}:
            raise serializers.ValidationError("检索字段或匹配方式错误")
        if params.get("q"):
            qs = qs.filter(**{field + ("__icontains" if match == "contains" else "__iexact"): params["q"]})
        total = qs.count()
        rows = list(qs.order_by("source_order", "id").values("code", "name", "specification", "site")[(page - 1) * size:page * size])
        return response.Response({"count": total, "results": rows, "page": page, "page_size": size})

    @decorators.action(detail=False, methods=["post"])
    def export(self, request):
        params = request.data.get("filters", {})
        if not isinstance(params, dict):
            raise serializers.ValidationError("查询条件格式错误")
        batch = batch_for(params)
        if batch is None:
            raise serializers.ValidationError("尚未同步报表数据")
        qs = filtered_rows(batch, self.report_kind, params)
        columns = INVENTORY_COLUMNS if self.report_kind == "inventory" else MOVEMENT_COLUMNS
        book = Workbook(write_only=True)
        info = book.create_sheet("查询条件")
        info.append(["数据来源", batch.source])
        info.append(["同步时间", str(batch.created_at)])
        info.append(["流水范围", str(batch.manifest.get("date_from", "")), str(batch.manifest.get("date_to", ""))])
        for key, value in params.items():
            info.append([safe_text(key), safe_text(value)])
        sheet = book.create_sheet("明细")
        sheet.append([label for _, label in columns])
        for row in qs.iterator(chunk_size=1000):
            values = []
            for key, _label in columns:
                value = row.payload.get(key)
                values.append(None if value is None else Decimal(str(value)) if key in NUMERIC_COLUMNS else safe_text(value))
            sheet.append(values)
        output = SpooledTemporaryFile(max_size=8 * 1024 * 1024)
        try:
            book.save(output)
            output.seek(0)
            return FileResponse(output, as_attachment=True, filename=("物料库存明细" if self.report_kind == "inventory" else "物料进出明细") + ".xlsx")
        except Exception:
            output.close()
            raise
