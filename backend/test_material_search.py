from io import BytesIO
from openpyxl import load_workbook
from rest_framework.test import APITestCase
from backend.test_support import CoreFlowSupport
from backend.models import Material
from backend.domain.inventory_analysis import InventoryAnalysisSnapshot
from backend.master_data.inventory_analysis_views import bulk_source, snapshot_payload


class MaterialSearchTests(CoreFlowSupport, APITestCase):
    def test_pagination_filters_candidates_and_export(self):
        ctx = self.base_context()
        original = Material.objects.get(pk=ctx['material']['id'])
        Material.objects.bulk_create([Material(code=f'TEST-{i:03}', name='检索测试',
            specification='分页', category=original.category, uom=original.uom,
            default_location=original.default_location, active=False) for i in range(105)])
        params = {'code': 'TEST-', 'q': '分页', 'category': original.category.code, 'active': '0'}
        url = '/api/material/search/'
        with self.assertNumQueries(2):
            first = self.client.get(url, params)
        self.assertEqual(first.status_code, 200, first.data)
        self.assertEqual(first.data['count'], 105)
        self.assertEqual(len(first.data['results']), 50)
        last = self.client.get(url, {**params, 'page': 3}).data
        self.assertEqual(len(last['results']), 5)
        self.assertEqual(last['results'][-1]['code'], 'TEST-104')
        self.assertEqual(self.client.get(url, {'code': 'TEST-001', 'exact': '1'}).data['count'], 1)
        self.assertEqual(self.client.get(url, {'page_size': 101}).status_code, 400)
        self.assertEqual(self.client.get(url, {'q': '不存在'}).data['count'], 0)
        for kind in ['categories', 'uoms', 'taxCodes', 'locations', 'suppliers']:
            r = self.client.get('/api/material/candidates/', {'kind': kind})
            self.assertEqual(r.status_code, 200, r.data)
        export = self.client.post('/api/material/search-export/', {'filters': params}, format='json')
        self.assertEqual(export.status_code, 200)
        sheet = load_workbook(BytesIO(export.content)).active
        self.assertEqual(sheet.max_row, 106)
        self.assertEqual(self.client.get(f'/api/material/{original.pk}/').data['code'], original.code)

    def test_snapshot_cache_changes_with_version(self):
        snapshot_payload.cache_clear()
        row = InventoryAnalysisSnapshot.objects.create(kind='bulk', site='1000', payload={'rows': [], 'details': []})
        self.assertEqual(bulk_source('1000', 'dgyzx1')[1]['rows'], [])
        row.payload = {'rows': [{'tmp_part': 'NEW'}], 'details': []}
        row.save()
        self.assertEqual(bulk_source('1000', 'dgyzx1')[1]['rows'], [{'tmp_part': 'NEW'}])
        snapshot_payload.cache_clear()
