"""Spreadsheet reading, preview validation and transactional material import."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from backend.material_codes import allocate_material_codes
from backend.models import Company, Location, Material, MaterialCompany, MaterialUomConversion, Partner, ProductCategory, TaxCode, Uom

from .material_mapping import REQUIRED_COLUMNS, clean_text, json_value, normalize_row, normalized_headers
from .models import ErpImportProfile, MaterialImportBatch, MaterialImportRow


def _read_rows(uploaded_file):
    suffix = Path(uploaded_file.name).suffix.lower()
    if suffix not in {".xls", ".xlsx"}:
        raise ValueError("请上传 .xls 或 .xlsx 格式的物料导入模板")
    if suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=uploaded_file.read())
        sheet = next((item for item in book.sheets() if item.name.lower() == "pt_mstr"), None)
        if sheet is None:
            raise ValueError("未找到 pt_mstr 工作表")
        data = [[sheet.cell_value(row, column) for column in range(sheet.ncols)] for row in range(sheet.nrows)]
    else:
        from openpyxl import load_workbook
        book = load_workbook(uploaded_file, read_only=True, data_only=True)
        sheet = book["pt_mstr"] if "pt_mstr" in book.sheetnames else None
        if sheet is None:
            raise ValueError("未找到 pt_mstr 工作表")
        data = [list(row) for row in sheet.iter_rows(values_only=True)]
    if len(data) < 2:
        raise ValueError("pt_mstr 工作表没有可导入的数据")
    headers = normalized_headers(data[0])
    missing = REQUIRED_COLUMNS - set(headers)
    if missing:
        raise ValueError(f"模板缺少必要字段：{'、'.join(sorted(missing))}")
    has_title_row = len(data) > 2 and clean_text(data[1][0] if data[1] else "") in {"物料编码", "pt_part"}
    return headers, data[2:] if has_title_row else data[1:], 3 if has_title_row else 2


def _validation_context():
    profile = ErpImportProfile.objects.select_related("company").filter(enabled=True, is_default=True).first()
    if profile is None:
        company = Company.objects.first() if Company.objects.count() == 1 else None
        if company is None:
            raise ValueError("请先配置唯一的启用默认导入主体")
        profile, _ = ErpImportProfile.objects.get_or_create(
            enabled=True,
            is_default=True,
            defaults={"name": "coerp_dgyzx1 单主体导入", "company": company},
        )
    return {
        "categories": {item.code.lower(): item for item in ProductCategory.objects.all()},
        "uoms": {value.lower(): item for item in Uom.objects.all() for value in (item.code, item.name)},
        "locations": {item.code.lower(): item for item in Location.objects.all()},
        "suppliers": {item.code.lower(): item for item in Partner.objects.filter(kind__in=(Partner.PartnerKind.SUPPLIER, Partner.PartnerKind.BOTH))},
        "tax_codes": {item.code.lower(): item for item in TaxCode.objects.filter(active=True)},
        "companies": {value.lower(): item for item in Company.objects.all() for value in (item.code, item.name)},
        "import_profile": profile,
        "existing_codes": set(Material.objects.values_list("code", flat=True)),
    }


def _validate(normalized, context, default_tax_code, seen_codes):
    errors = []
    category = context["categories"].get(normalized["category_code"].lower())
    uom = context["uoms"].get(normalized["uom_code"].lower())
    location = context["locations"].get(normalized["location_code"].lower())
    supplier = context["suppliers"].get(normalized["supplier_code"].lower()) if normalized["supplier_code"] else None
    tax_code = normalized["tax_code"] or default_tax_code
    tax_master = context["tax_codes"].get(tax_code.lower()) if tax_code else None
    if not normalized["name"]:
        errors.append("物料名称不能为空")
    if not category:
        errors.append(f"产品类不存在：{normalized['category_code'] or '空'}")
    if not uom:
        errors.append(f"计量单位不存在：{normalized['uom_code'] or '空'}")
    if not location:
        errors.append(f"库位不存在：{normalized['location_code'] or '空'}")
    if normalized["supplier_code"] and not supplier:
        errors.append(f"供应商不存在或不是供应商：{normalized['supplier_code']}")
    if not tax_code:
        errors.append("税务编码不能为空，请在模板或预检页面提供默认税务编码")
    elif context["tax_codes"] and not tax_master:
        errors.append(f"税务编码不存在或已停用：{tax_code}")
    source_code = normalized["source_code"]
    if source_code and (source_code in context["existing_codes"] or source_code in seen_codes):
        errors.append(f"物料编码已存在：{source_code}")
    if source_code:
        seen_codes.add(source_code)
    profile = context["import_profile"]
    selected = [column for column, enabled in normalized["company_columns"].items() if enabled]
    normalized["company_selection"] = {"mode": "defaulted", "column": "", "company_code": "", "warnings": []}
    if not profile:
        errors.append("未配置物料导入主体公司，请先配置 ERP 导入绑定")
    else:
        current_column = profile.source_company_column.lower()
        if len(selected) > 1:
            errors.append("同一行不能选择多个主体公司")
            normalized["company_selection"]["mode"] = "invalid"
        elif selected and selected[0] != current_column:
            errors.append(f"首期仅支持主体列 {profile.source_company_column}，当前选择：{selected[0]}")
            normalized["company_selection"]["mode"] = "invalid"
        else:
            normalized["company_selection"].update({
                "mode": "explicit" if selected else "defaulted",
                "column": current_column,
                "company_code": profile.company.code,
                "warnings": [] if selected else [f"未填写 {profile.source_company_column}，已自动归属当前主体 {profile.company.name}"],
            })
    if normalized["batch_control"] and not normalized["batch_rule"]:
        errors.append("启用批号控制时必须填写批号规则")
    for conversion in normalized["conversions"]:
        if conversion["uom_code"].lower() not in context["uoms"]:
            errors.append(f"业务单位不存在：{conversion['uom_code']}")
        if conversion["business_qty"] == "0" or conversion["stock_qty"] == "0":
            errors.append(f"业务单位换算系数必须大于 0：{conversion['uom_code']}")
    normalized["tax_code"] = tax_code
    return errors


def preview_material_import(uploaded_file, user, default_tax_code=""):
    headers, spreadsheet_rows, first_row_number = _read_rows(uploaded_file)
    rows = []
    for offset, cells in enumerate(spreadsheet_rows):
        raw = {headers[index]: json_value(cells[index]) if index < len(cells) else "" for index in range(len(headers))}
        if not any(clean_text(value) for value in raw.values()):
            continue
        rows.append((first_row_number + offset, raw, normalize_row(raw)))
    context = _validation_context()
    seen_codes = set()
    checked = []
    for number, raw, normalized in rows:
        errors = _validate(normalized, context, default_tax_code.strip(), seen_codes)
        checked.append((number, raw, normalized, errors))
    pending_by_prefix = defaultdict(list)
    for index, (_, _, normalized, errors) in enumerate(checked):
        if not errors and not normalized["source_code"]:
            category = context["categories"][normalized["category_code"].lower()]
            pending_by_prefix[(category.code_prefix or category.code).strip()].append(index)
    allocated = {}
    for prefix, indexes in pending_by_prefix.items():
        local_codes = Material.objects.filter(code__startswith=f"{prefix}-").values_list("code", flat=True)
        allocated.update(zip(indexes, allocate_material_codes(prefix, local_codes, len(indexes))))

    batch = MaterialImportBatch.objects.create(
        source_file=uploaded_file, original_filename=uploaded_file.name, default_tax_code=default_tax_code.strip(),
        total_rows=len(checked), valid_rows=sum(not errors for *_, errors in checked),
        error_rows=sum(bool(errors) for *_, errors in checked), created_by=user.get_username(),
    )
    MaterialImportRow.objects.bulk_create([
        MaterialImportRow(batch=batch, row_number=number, source_data=raw, normalized_data=normalized,
                          allocated_code=normalized["source_code"] or allocated.get(index, ""),
                          status=MaterialImportRow.Status.ERROR if errors else MaterialImportRow.Status.VALID,
                          errors=errors)
        for index, (number, raw, normalized, errors) in enumerate(checked)
    ])
    return batch


def serialize_batch(batch, include_rows=True):
    data = {
        "id": batch.id, "original_filename": batch.original_filename, "status": batch.status,
        "total_rows": batch.total_rows, "valid_rows": batch.valid_rows, "error_rows": batch.error_rows,
        "imported_rows": batch.imported_rows, "created_by": batch.created_by, "created_at": batch.created_at,
        "confirmed_by": batch.confirmed_by, "confirmed_at": batch.confirmed_at,
    }
    profile = ErpImportProfile.objects.select_related("company").filter(enabled=True, is_default=True).first()
    data["import_profile"] = {
        "name": profile.name,
        "company_code": profile.company.code,
        "company_name": profile.company.name,
        "source_database": profile.source_database,
        "source_company_column": profile.source_company_column,
    } if profile else None
    if include_rows:
        data["rows"] = [{"id": row.id, "row_number": row.row_number, "allocated_code": row.allocated_code,
                         "status": row.status, "errors": row.errors, "name": row.normalized_data.get("name", ""),
                         "category_code": row.normalized_data.get("category_code", ""), "material_id": row.material_id,
                         "company_selection": row.normalized_data.get("company_selection", {})}
                        for row in batch.rows.all()]
    return data


@transaction.atomic
def confirm_material_import(batch, user):
    if batch.status != MaterialImportBatch.Status.PREVIEW:
        raise ValueError("该批次已经确认，不能重复导入")
    if batch.error_rows:
        raise ValueError("预检仍有错误行，修正后请重新上传")
    rows = list(batch.rows.select_for_update().order_by("row_number"))
    context = _validation_context()
    codes = set()
    for row in rows:
        data = row.normalized_data
        errors = _validate(data, context, batch.default_tax_code, codes)
        if errors:
            row.errors = errors
            row.status = MaterialImportRow.Status.ERROR
            row.save(update_fields=("errors", "status"))
            batch.error_rows += 1
    if batch.error_rows:
        batch.valid_rows = len(rows) - batch.error_rows
        batch.save(update_fields=("valid_rows", "error_rows"))
        raise ValueError("基础资料已变化，预检不再通过，请重新上传")
    for row in rows:
        data = row.normalized_data
        category = context["categories"][data["category_code"].lower()]
        # Lock categories before saving to keep an allocated sequence consistent.
        category = ProductCategory.objects.select_for_update().get(pk=category.pk)
        material = Material(code=row.allocated_code, category=category, uom=context["uoms"][data["uom_code"].lower()],
                            default_location=context["locations"][data["location_code"].lower()],
                            default_supplier=context["suppliers"].get(data["supplier_code"].lower()),
                            tax_master=context["tax_codes"].get(data["tax_code"].lower()), created_by=user.get_username(),
                            tax_code=data["tax_code"],
                            **{key: value for key, value in data.items() if key not in {"source_code", "category_code", "uom_code", "location_code", "supplier_code", "tax_code", "company_columns", "company_selection", "conversions"}})
        try:
            material.full_clean()
        except ValidationError as exc:
            raise ValueError("第 %s 行无法导入：%s" % (row.row_number, exc.messages)) from exc
        material.save()
        MaterialCompany.objects.create(material=material, company=context["import_profile"].company)
        for conversion in data["conversions"]:
            business_uom = context["uoms"].get(conversion["uom_code"].lower())
            if not business_uom:
                raise ValueError(f"第 {row.row_number} 行的业务单位不存在：{conversion['uom_code']}")
            if conversion["business_qty"] == "0" or conversion["stock_qty"] == "0":
                continue
            item = MaterialUomConversion(material=material, usage=conversion["usage"], business_uom=business_uom,
                                         business_qty=conversion["business_qty"], stock_qty=conversion["stock_qty"])
            item.full_clean()
            item.save()
        row.material = material
        row.status = MaterialImportRow.Status.IMPORTED
        row.save(update_fields=("material", "status"))
    batch.status = MaterialImportBatch.Status.IMPORTED
    batch.imported_rows = len(rows)
    batch.confirmed_by = user.get_username()
    batch.confirmed_at = timezone.now()
    batch.save(update_fields=("status", "imported_rows", "confirmed_by", "confirmed_at"))
    return batch
