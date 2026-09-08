from io import BytesIO

from django.contrib.auth import get_user_model
from django.urls import resolve
from rest_framework.test import APITestCase
from openpyxl import load_workbook

from backend.models import ApprovalStatus, BusinessGroup, Department, Role, UserRole
from backend.test_support import SourceDatabaseIsolatedMixin


class MasterDataApiTests(SourceDatabaseIsolatedMixin, APITestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser("erp-admin", password="test-pass")
        self.client.force_authenticate(self.user)

    def post(self, resource, payload):
        response = self.client.post(f"/api/{resource}/", payload, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        return response.data

    def test_material_uses_master_data_and_audited_approval_workflow(self):
        uom_category = self.post("uom-category", {"code": "COUNT", "name": "数量"})
        uom = self.post("uom", {"code": "JIAN", "name": "件", "category": uom_category["id"]})
        location = self.post("location", {
            "code": "FG01", "name": "成品仓", "location_type": "fg", "location_attribute": "成品存放",
        })
        self.assertEqual(location["location_attribute"], "成品存放")
        category = self.post("product-category", {
            "code": "FG", "name": "成品", "default_uom": uom["id"], "default_location": location["id"],
        })
        company = self.post("company", {"code": "C001", "name": "合越制造"})
        material = self.post("material", {
            "name": "注塑外壳", "category": category["id"], "uom": uom["id"],
            "default_location": location["id"], "tax_code": "3926909090",
        })
        self.assertEqual(material["code"], "FG-0001")
        self.assertEqual(material["default_location"], location["id"])
        self.assertEqual(material["created_by"], self.user.username)
        self.assertTrue(material["created_at"])
        company_scope = self.post("material-company", {"material": material["id"], "company": company["id"], "enabled": True})
        self.assertEqual(company_scope["company"], company["id"])

        submitted = self.client.post(f"/api/material/{material['id']}/submit/", {}, format="json")
        self.assertEqual(submitted.status_code, 200, submitted.data)
        self.assertEqual(submitted.data["status"], "pending")
        approved = self.client.post(f"/api/material/{material['id']}/approve/", {"actor": "工程审核"}, format="json")
        self.assertEqual(approved.status_code, 200, approved.data)
        self.assertEqual(approved.data["status"], "approved")
        self.assertEqual(approved.data["approved_by"], self.user.username)
        blocked = self.client.patch(f"/api/material/{material['id']}/", {"name": "不应修改"}, format="json")
        self.assertEqual(blocked.status_code, 400, blocked.data)
        unapproved = self.client.post(f"/api/material/{material['id']}/unapprove/", {}, format="json")
        self.assertEqual(unapproved.status_code, 200, unapproved.data)
        self.assertEqual(unapproved.data["status"], "draft")
        forged_status = self.client.patch(
            f"/api/material/{material['id']}/", {"status": "approved"}, format="json",
        )
        self.assertEqual(forged_status.status_code, 400, forged_status.data)
        editable = self.client.patch(f"/api/material/{material['id']}/", {"name": "反审核后可修改"}, format="json")
        self.assertEqual(editable.status_code, 200, editable.data)

    def test_audit_fields_are_set_from_the_authenticated_user_and_cannot_be_overridden(self):
        currency = self.post("currency", {
            "code": "CNY", "name": "人民币", "created_by": "伪造用户",
        })
        self.assertEqual(currency["created_by"], self.user.username)
        self.assertTrue(currency["created_at"])

        updated = self.client.patch(f"/api/currency/{currency['id']}/", {
            "name": "人民币元", "updated_by": "伪造用户",
        }, format="json")
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual(updated.data["updated_by"], self.user.username)
        self.assertTrue(updated.data["updated_at"])

    def test_supplier_quote_requires_supplier_and_rejects_overlapping_dates(self):
        business_group = BusinessGroup.objects.get(code="01")
        uom_category = self.post("uom-category", {"code": "WEIGHT", "name": "重量"})
        uom = self.post("uom", {"code": "QIANKE", "name": "千克", "category": uom_category["id"]})
        location = self.post("location", {"code": "RM01", "name": "原料仓", "location_type": "rm"})
        category = self.post("product-category", {"code": "RM", "name": "原材料"})
        currency = self.post("currency", {"code": "CNY", "name": "人民币", "symbol": "¥"})
        supplier = self.post("partner", {
            "code": "SUP001", "name": "原料供应商", "short_name": "原料供", "kind": "supplier", "payment_method": "供应商月结", "currency": currency["id"],
        })
        material = self.post("material", {
            "name": "聚乙烯颗粒", "category": category["id"], "uom": uom["id"],
            "default_location": location["id"], "tax_code": "3901200090",
        })
        for resource, record_id in (("partner", supplier["id"]), ("material", material["id"])):
            submitted = self.client.post(f"/api/{resource}/{record_id}/submit/", {}, format="json")
            self.assertEqual(submitted.status_code, 200, submitted.data)
            approved = self.client.post(f"/api/{resource}/{record_id}/approve/", {"actor": "审核人"}, format="json")
            self.assertEqual(approved.status_code, 200, approved.data)
        quote = self.post("supplier-quote", {
            "supplier": supplier["id"], "material": material["id"], "business_group": business_group.id, "currency": currency["id"],
            "quote_type": "purchase", "tax_included": True, "tax_rate": "13", "effective_date": "2026-01-01",
            "expiry_date": "2026-06-30", "material_unit_price": "10.25",
        })
        self.assertTrue(quote["number"].startswith("MYPQ"))
        overlap = self.client.post("/api/supplier-quote/", {
            "supplier": supplier["id"], "material": material["id"], "business_group": business_group.id, "currency": currency["id"],
            "quote_type": "purchase", "effective_date": "2026-06-01", "expiry_date": "2026-12-31",
            "material_unit_price": "10.50",
        }, format="json")
        self.assertEqual(overlap.status_code, 400)
        self.assertIn("有效日期", str(overlap.data))

    def test_supplier_quote_persists_business_group(self):
        category = self.post("product-category", {"code": "BG", "name": "业务组报价测试"})
        uom_category = self.post("uom-category", {"code": "BGCOUNT", "name": "业务组数量"})
        uom = self.post("uom", {"code": "BGJIAN", "name": "业务组件", "category": uom_category["id"]})
        location = self.post("location", {"code": "BGRM", "name": "业务组仓", "location_type": "rm"})
        currency = self.post("currency", {"code": "BGCNY", "name": "业务组人民币"})
        groups = self.client.get("/api/business-group/")
        self.assertEqual(groups.status_code, 200, groups.data)
        group = next(group for group in groups.data if group["code"] == "01")
        supplier = self.post("partner", {
            "code": "BGSUP", "name": "业务组供应商", "short_name": "业务组供", "kind": "supplier",
            "payment_method": "月结", "currency": currency["id"],
        })
        material = self.post("material", {
            "name": "业务组物料", "category": category["id"], "uom": uom["id"],
            "default_location": location["id"], "tax_code": "BG-TAX",
        })
        for resource, record_id in (("partner", supplier["id"]), ("material", material["id"])):
            self.client.post(f"/api/{resource}/{record_id}/submit/", {}, format="json")
            self.client.post(f"/api/{resource}/{record_id}/approve/", {}, format="json")
        quote = self.post("supplier-quote", {
            "supplier": supplier["id"], "material": material["id"], "business_group": group["id"],
            "quote_type": "purchase", "effective_date": "2026-08-01", "material_unit_price": "10",
        })
        self.assertEqual(quote["business_group"], group["id"])
        self.assertEqual(quote["business_group_code"], "01")
        self.assertEqual(quote["business_group_name"], "业务一部")
        self.assertTrue(BusinessGroup.objects.filter(code="01", supplier_quotes__pk=quote["id"]).exists())

    def test_supplier_quote_tiers_workflow_and_price_resolution(self):
        business_group = BusinessGroup.objects.get(code="01")
        category = self.post("product-category", {"code": "SQ", "name": "供应商报价测试"})
        uom_category = self.post("uom-category", {"code": "SQCOUNT", "name": "报价数量"})
        uom = self.post("uom", {"code": "SQJIAN", "name": "报价件", "category": uom_category["id"]})
        location = self.post("location", {"code": "SQRM", "name": "报价仓", "location_type": "rm"})
        currency = self.post("currency", {"code": "SQCNY", "name": "报价人民币"})
        supplier = self.post("partner", {
            "code": "SQSUP", "name": "报价供应商", "short_name": "报价供", "kind": "supplier",
            "payment_method": "月结", "currency": currency["id"], "tax_rate": "13",
        })
        material = self.post("material", {
            "name": "报价物料", "category": category["id"], "uom": uom["id"],
            "default_location": location["id"], "tax_code": "SQ-TAX", "min_purchase_qty": "5", "min_pack_qty": "10",
        })
        for resource, record_id in (("partner", supplier["id"]), ("material", material["id"])):
            self.client.post(f"/api/{resource}/{record_id}/submit/", {}, format="json")
            self.client.post(f"/api/{resource}/{record_id}/approve/", {}, format="json")
        quote = self.post("supplier-quote", {
            "supplier": supplier["id"], "material": material["id"], "business_group": business_group.id, "quote_type": "purchase",
            "effective_date": "2026-08-01", "material_unit_price": "20", "lines": [
                {"line_number": 10, "min_qty": "100", "material_unit_price": "18", "notes": "百件价"},
                {"line_number": 20, "min_qty": "500", "material_unit_price": "16", "notes": "五百件价"},
            ],
        })
        self.assertTrue(quote["number"].startswith("MYPQ"))
        self.assertEqual(quote["currency"], currency["id"])
        self.assertEqual(quote["tax_rate"], "13.00000000")
        self.assertEqual(quote["effective_min_purchase_qty"], "5.00000000")
        self.assertEqual(quote["effective_min_pack_qty"], "10.00000000")
        self.assertEqual(self.client.post(f"/api/supplier-quote/{quote['id']}/confirm/", {}, format="json").status_code, 200)
        self.assertEqual(self.client.post(f"/api/supplier-quote/{quote['id']}/approve/", {}, format="json").status_code, 200)
        self.assertEqual(self.client.post(f"/api/supplier-quote/{quote['id']}/ratify/", {}, format="json").status_code, 200)
        resolved = self.client.get(f"/api/supplier-quote/resolve/?supplier={supplier['id']}&material={material['id']}&quantity=520&date=2026-08-15")
        self.assertEqual(resolved.status_code, 200, resolved.data)
        self.assertEqual(resolved.data["resolved_unit_price"], "16.00000000")
        base_price = self.client.get(f"/api/supplier-quote/resolve/?supplier={supplier['id']}&material={material['id']}&quantity=50&date=2026-08-15")
        self.assertEqual(base_price.status_code, 200, base_price.data)
        self.assertEqual(base_price.data["resolved_unit_price"], "20.00000000")
        sales_confirmation = self.client.post(f"/api/supplier-quote/{quote['id']}/sales-confirm/", {}, format="json")
        self.assertEqual(sales_confirmation.status_code, 200, sales_confirmation.data)
        self.assertTrue(sales_confirmation.data["is_sales_confirmed"])
        exported = self.client.post("/api/supplier-quote/export/", {"ids": [quote["id"]]}, format="json")
        self.assertEqual(exported.status_code, 200)
        workbook = load_workbook(BytesIO(exported.content), read_only=True)
        self.assertEqual(workbook.sheetnames, ["6.5 -- 供应商报价(PUVQMTA1)", "报价阶梯明细"])
        self.assertEqual(sum(1 for _ in workbook["报价阶梯明细"].iter_rows()), 3)

    def test_partner_customer_supplier_alignment_fields_are_persisted(self):
        customer = self.post("partner", {
            "code": "CUS001", "name": "客户公司", "short_name": "客户", "kind": "customer",
            "payment_method": "月结", "notes": "客户备注", "copper_origin": "华东", "copper_currency": "CNY",
            "round_order_quantity_for_tier_price": True,
        })
        self.assertEqual(customer["notes"], "客户备注")
        self.assertEqual(customer["copper_origin"], "华东")
        self.assertTrue(customer["round_order_quantity_for_tier_price"])

        supplier = self.post("partner", {
            "code": "SUP003", "name": "供应商公司", "short_name": "供应商", "kind": "supplier",
            "payment_method": "供应商月结", "buyer": "采购员", "notes": "供应商备注", "internal_company": True,
        })
        self.assertEqual(supplier["buyer"], "采购员")
        self.assertEqual(supplier["notes"], "供应商备注")
        self.assertTrue(supplier["internal_company"])

    def test_customer_contacts_are_saved_with_the_customer_record(self):
        customer = self.post("partner", {
            "code": "CUSCONTACT", "name": "联系人客户", "short_name": "联系人",
            "kind": "customer", "contacts": [
                {"name": "张三", "position": "采购", "phone": "13800000000", "fax": "0755-1000", "email": "zhang@example.com", "is_primary": True},
                {"name": "李四", "position": "工程", "phone": "13900000000", "email": "li@example.com", "is_primary": False},
            ],
        })
        self.assertEqual(customer["created_by"], self.user.username)
        self.assertEqual([contact["name"] for contact in customer["contacts"]], ["张三", "李四"])

        updated = self.client.patch(f"/api/partner/{customer['id']}/", {
            "contacts": [{"name": "王五", "position": "财务", "phone": "13700000000", "email": "wang@example.com", "is_primary": True}],
        }, format="json")
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual([contact["name"] for contact in updated.data["contacts"]], ["王五"])

    def test_customer_material_follows_customer_material_code_composite_key_and_audits_actor(self):
        uom_category = self.post("uom-category", {"code": "CMCOUNT", "name": "数量"})
        uom = self.post("uom", {"code": "CMJIAN", "name": "件", "category": uom_category["id"]})
        location = self.post("location", {"code": "CMFG", "name": "客户物料仓", "location_type": "fg"})
        category = self.post("product-category", {"code": "CM", "name": "客户物料类"})
        customer = self.post("partner", {"code": "CUSMAT", "name": "客户物料客户", "short_name": "客户物料", "kind": "customer"})
        supplier = self.post("partner", {"code": "SUPMAT", "name": "客户物料供应商", "short_name": "物料供应", "kind": "supplier", "payment_method": "月结"})
        materials = [self.post("material", {
            "name": name, "category": category["id"], "uom": uom["id"],
            "default_location": location["id"], "tax_code": f"TAX{index}",
        }) for index, name in enumerate(("本方物料甲", "本方物料乙"), start=1)]

        payload = {
            "customer": customer["id"], "material": materials[0]["id"], "customer_code": "PORT-001",
            "customer_name": "客户端物料", "customer_uom": uom["id"], "customer_uom_rate_m": "2",
            "customer_uom_rate_d": "1", "terminal_customer_code": "END-001", "terminal_customer_name": "终端客户",
            "created_by": "伪造账号",
        }
        created = self.post("customer-material", payload)
        self.assertEqual(created["material_code"], materials[0]["code"])
        self.assertEqual(created["material_name"], "本方物料甲")
        self.assertEqual(created["material_specification"], "")
        self.assertEqual(created["created_by"], self.user.username)
        self.assertTrue(created["created_at"])

        duplicate = self.client.post("/api/customer-material/", payload, format="json")
        self.assertEqual(duplicate.status_code, 400)
        second = self.post("customer-material", {**payload, "material": materials[1]["id"]})
        self.assertEqual(second["customer_code"], "PORT-001")

        not_a_customer = self.client.post("/api/customer-material/", {**payload, "customer": supplier["id"], "customer_code": "PORT-002"}, format="json")
        self.assertEqual(not_a_customer.status_code, 400)
        self.assertIn("只能关联客户", str(not_a_customer.data))

    def test_master_data_exports_are_compatible_xlsx_files(self):
        department = self.post("department", {"code": "EXP-DP", "name": "导出部门"})
        currency = self.post("currency", {"code": "CNY", "name": "人民币", "symbol": "¥"})
        uom_category = self.post("uom-category", {"code": "EXP-COUNT", "name": "导出数量"})
        uom = self.post("uom", {"code": "EXP-JIAN", "name": "导出件", "category": uom_category["id"]})
        location = self.post("location", {
            "code": "EXP-FG", "name": "导出仓", "location_type": "fg",
            "quarantine_return": True, "shipped_goods_account": "1406", "sales_cost_account": "6401",
        })
        category = self.post("product-category", {
            "code": "EXP", "name": "导出产品类", "default_uom": uom["id"], "default_location": location["id"],
        })
        customer = self.post("partner", {
            "code": "EXP-CUS", "name": "导出客户有限公司", "short_name": "导出客户", "kind": "customer",
            "currency": currency["id"], "contacts": [{"name": "联系人甲", "position": "采购", "phone": "13800000000"}],
        })
        self.post("partner", {
            "code": "EXP-SUP", "name": "导出供应商有限公司", "short_name": "导出供应商", "kind": "supplier",
            "currency": currency["id"], "payment_method": "月结30天",
        })
        material = self.post("material", {
            "name": "=导出物料", "category": category["id"], "uom": uom["id"],
            "default_location": location["id"], "tax_code": "EXP-TAX", "specification": "A-01",
        })
        customer_material = self.post("customer-material", {
            "customer": customer["id"], "material": material["id"], "customer_code": "CUS-PART-01",
            "customer_name": "客户物料一", "customer_uom": uom["id"], "customer_uom_rate_m": "2",
            "customer_uom_rate_d": "1", "terminal_customer_code": "END-001", "terminal_customer_name": "终端客户一",
        })

        resources = {
            "department": ({"ids": [department["id"]]}, "部门代码"),
            "currency": ({"ids": [currency["id"]]}, "货币"),
            "location": ({"ids": [location["id"]]}, "检验不良退货暂存仓"),
            "product-category": ({"ids": [category["id"]]}, "产品类"),
            "uom": ({"ids": [uom["id"]]}, "单位"),
            "material": ({"ids": [material["id"]]}, "物料编码"),
            "partner-customer": ({"ids": [customer["id"]], "partner_kind": "customer"}, "客户代码"),
            "partner-supplier": ({"partner_kind": "supplier"}, "供应商代码"),
            "customer-material": ({"ids": [customer_material["id"]]}, "客户物料编码"),
        }
        workbooks = {}
        for key, (payload, expected_header) in resources.items():
            resource = "partner" if key.startswith("partner-") else key
            response = self.client.post(f"/api/{resource}/export/", payload, format="json")
            self.assertEqual(response.status_code, 200, getattr(response, "data", None))
            self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            workbook = load_workbook(BytesIO(response.content), read_only=False, data_only=False)
            sheet = workbook.active
            headers = [cell.value for cell in sheet[1]]
            self.assertIn(expected_header, headers)
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertTrue(sheet.auto_filter.ref)
            self.assertFalse(sheet.sheet_view.showGridLines)
            self.assertEqual(sheet.row_dimensions[1].height, 28)
            self.assertEqual(sheet.row_dimensions[2].height, 20)
            self.assertIsNotNone(sheet.cell(1, 1).border.left.style)
            self.assertIsNotNone(sheet.cell(2, 1).border.left.style)
            self.assertEqual(sheet.max_row, 2)
            workbooks[key] = (sheet, headers)

        customer_sheet, customer_headers = workbooks["partner-customer"]
        self.assertEqual(customer_sheet.cell(2, customer_headers.index("联系人1") + 1).value, "联系人甲")
        self.assertEqual(customer_sheet.cell(2, customer_headers.index("职位1") + 1).value, "采购")
        material_sheet, material_headers = workbooks["material"]
        self.assertEqual(material_sheet.cell(2, material_headers.index("物料名称") + 1).value, "'=导出物料")
        mapping_sheet, mapping_headers = workbooks["customer-material"]
        self.assertEqual(mapping_sheet.cell(2, mapping_headers.index("物料名称") + 1).value, "'=导出物料")
        self.assertEqual(mapping_sheet.cell(2, mapping_headers.index("物料规格") + 1).value, "A-01")
        self.assertEqual(mapping_sheet.cell(2, mapping_headers.index("终端客户编码") + 1).value, "END-001")
        self.assertEqual(mapping_sheet.cell(2, mapping_headers.index("录入人") + 1).value, self.user.username)
        self.assertGreaterEqual(material_sheet.column_dimensions["A"].width, 18)
        self.assertGreaterEqual(mapping_sheet.column_dimensions["C"].width, 18)

    def test_export_rejects_invalid_filter_ids_and_requires_partner_kind(self):
        invalid_ids = self.client.post("/api/material/export/", {"ids": ["1"]}, format="json")
        self.assertEqual(invalid_ids.status_code, 400)
        missing_kind = self.client.post("/api/partner/export/", {}, format="json")
        self.assertEqual(missing_kind.status_code, 400)

    def test_export_requires_the_resource_view_permission(self):
        department = Department.objects.create(code="PERM-DP", name="权限部门", status=ApprovalStatus.APPROVED)
        role = Role.objects.create(
            code="EXPORT-VIEW", name="导出查看", department=department,
            permissions=["department.view"], status=ApprovalStatus.APPROVED,
        )
        user = get_user_model().objects.create_user("export-viewer", password="test-pass")
        UserRole.objects.create(user=user, role=role, department=department, active=True)
        self.client.force_authenticate(user)

        allowed = self.client.post("/api/department/export/", {}, format="json")
        self.assertEqual(allowed.status_code, 200)
        role.permissions = []
        role.save(update_fields=("permissions",))
        denied = self.client.post("/api/department/export/", {}, format="json")
        self.assertEqual(denied.status_code, 403)

    def test_quote_rejects_unapproved_supplier(self):
        uom_category = self.post("uom-category", {"code": "LENGTH", "name": "长度"})
        uom = self.post("uom", {"code": "MI", "name": "米", "category": uom_category["id"]})
        location = self.post("location", {"code": "RM02", "name": "材料仓", "location_type": "rm"})
        category = self.post("product-category", {"code": "PK", "name": "包装类"})
        currency = self.post("currency", {"code": "USD", "name": "美元", "symbol": "$"})
        supplier = self.post("partner", {
            "code": "SUP002", "name": "待审供应商", "short_name": "待审供", "kind": "supplier", "payment_method": "供应商月结",
        })
        material = self.post("material", {
            "name": "包装膜", "category": category["id"], "uom": uom["id"],
            "default_location": location["id"], "tax_code": "3920109090",
        })
        response = self.client.post("/api/supplier-quote/", {
            "supplier": supplier["id"], "material": material["id"], "currency": currency["id"],
            "quote_type": "purchase", "effective_date": "2026-01-01", "material_unit_price": "2.5",
        }, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("供应商审核通过", str(response.data))

    def test_unauthenticated_requests_are_rejected(self):
        self.client.force_authenticate(user=None)
        response = self.client.get("/api/material/")
        self.assertIn(response.status_code, {401, 403})

    def test_login_returns_a_token_that_authenticates_api_requests(self):
        self.client.force_authenticate(user=None)
        login = self.client.post("/api/auth/login/", {
            "username": "erp-admin", "password": "test-pass",
        }, format="json")

        self.assertEqual(login.status_code, 200, login.data)
        self.assertEqual(login.data["username"], "erp-admin")
        self.assertTrue(login.data["token"])
        self.assertEqual(login.data["permissions"], ["*"])

        self.client.credentials(HTTP_AUTHORIZATION=f"Token {login.data['token']}")
        self.assertEqual(self.client.get("/api/material/").status_code, 200)

    def test_login_rejects_invalid_password(self):
        self.client.force_authenticate(user=None)
        response = self.client.post("/api/auth/login/", {
            "username": "erp-admin", "password": "wrong-password",
        }, format="json")

        self.assertEqual(response.status_code, 400)
        self.assertIn("用户名或密码", str(response.data))

    def test_domain_view_modules_keep_existing_api_routes(self):
        expected_modules = {
            "/api/location/": "backend.master_data.api_common",
            "/api/sales-quote/": "backend.master_data.sales_views",
            "/api/purchase-requisition/": "backend.master_data.purchase_views",
            "/api/inventory-alert/": "backend.master_data.inventory_views",
        }
        for path, module in expected_modules.items():
            self.assertEqual(resolve(path).func.cls.__module__, module)
