import json
from decimal import Decimal
from io import BytesIO
from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from rest_framework.test import APITestCase
from backend.domain.inventory_analysis import InventoryAnalysisSnapshot
from backend.inventory_analysis import aggregate_monthly, reconcile_bulk


class InventoryAnalysisTests(APITestCase):
    def setUp(self):
        self.client.force_authenticate(get_user_model().objects.create_superuser('analysis', password='test'))
        self.masters = {'materials': {'XSM-000832': {'pt_desc1': '铝模', 'pt_um': '个'}},
            'locations': {'R01': {'loc_desc': '原料仓', 'loc_ac_code_inv': ''}},
            'categories': {'XSM': {'pl_desc': '吸塑模具', 'pl_ac_code_inv': '1405'}},
            'suppliers': [], 'accounts': {'1405': '库存商品'}, 'groups': {}, 'part_types': {}, 'cost_types': {'RAW': '材料及成品库存仓'}}

    def monthly(self, **values):
        return dict(inb_part='XSM-000832', inb_loc='R01', inb_lot='', pt_prod_line='XSM', pt_group='', pt_part_type='', loc_type='原料仓', loc_cost_type='RAW', **values)

    def test_february_receipt_carries_to_march_opening_without_replaying_receipt(self):
        for period, values in [('2026-02', dict(inb_qty_in='1000', inb_qty_beg='0', inb_cost_in='10000', inb_cost_beg='0', inb_avg_cost='10', inb_120_qty='1000', inb_120_amt='10000')),
                               ('2026-03', dict(inb_qty_in='0', inb_qty_beg='1000', inb_cost_in='0', inb_cost_beg='10000', inb_avg_cost='0'))]:
            InventoryAnalysisSnapshot.objects.create(kind='monthly', site='1000', period=period,
                payload={'rows': [self.monthly(inb_qty_end='1000', inb_cost_end='10000', **values)], 'masters': self.masters})
        feb = self.client.get('/api/material-monthly-inventory/', {'period': '2026-02'}).data['results'][0]
        march = self.client.get('/api/material-monthly-inventory/', {'period': '2026-03'}).data['results'][0]
        self.assertEqual(feb['inb_qty_end'], march['inb_qty_beg'])
        self.assertEqual(feb['inb_cost_end'], march['inb_cost_beg'])
        self.assertEqual(march['inb_qty_in'], '0')
        self.assertEqual(march['unit_beg'], '10')
        self.assertEqual(feb['ac_code'], '1405')
        self.assertEqual(march['inb_avg_cost'], '0')

    def test_aggregation_uses_sum_ratio_but_client_average_for_actual_cost(self):
        rows = [self.monthly(inb_qty_end='2', inb_cost_end='10', inb_avg_cost='5'), self.monthly(inb_qty_end='8', inb_cost_end='80', inb_avg_cost='10')]
        result = aggregate_monthly(rows, ['inb_part', 'loc_cost_type'], self.masters)[0]
        self.assertEqual(result['unit_end'], '9')
        self.assertEqual(result['inb_avg_cost'], '7.5')
        self.assertNotIn('inb_loc', result)
        self.assertEqual(result['loc_desc'], '')

    def test_numeric_filter_global_sort_total_and_export(self):
        rows = [{'tmp_part': f'P{i:03}', 'pt_desc1': '=1+1', 'tmp_qty_remain': str(i-102)} for i in range(205)]
        InventoryAnalysisSnapshot.objects.create(kind='bulk', site='1000', payload={'rows': rows})
        url = '/api/material-inventory-batch/'
        filters = {'numeric_filters': json.dumps({'tmp_qty_remain': {'op': 'lt', 'value': '0'}}), 'sort': 'tmp_qty_remain'}
        result = self.client.get(url, filters).data
        self.assertEqual(result['count'], 102)
        self.assertEqual(len(result['results']), 100)
        self.assertEqual(result['results'][0]['tmp_qty_remain'], '-102')
        self.assertEqual(Decimal(result['totals']['tmp_qty_remain']), -5253)
        self.assertEqual(self.client.get(url, {**filters, 'direction': 'desc'}).data['results'][0]['tmp_qty_remain'], '-1')
        exported = self.client.post(url+'export/', {'filters': filters}, format='json')
        book = load_workbook(BytesIO(b''.join(exported.streaming_content)))
        self.assertEqual(book.worksheets[0].max_row, 103)
        self.assertNotEqual(book.worksheets[0]['B2'].data_type, 'f')
        for value in ('NaN', 'Infinity', 'bad'):
            invalid = {'numeric_filters': json.dumps({'tmp_qty_remain': {'op': 'eq', 'value': value}})}
            self.assertEqual(self.client.get(url, invalid).status_code, 400)
        self.assertEqual(self.client.get(url, {'sort': 'sql'}).status_code, 400)
        self.assertEqual(self.client.get('/api/material-monthly-inventory/', {'period': '1999-01'}).data['snapshot'], None)

    def test_permissions_and_wip_details(self):
        InventoryAnalysisSnapshot.objects.create(kind='bulk', site='1000', payload={'rows': [], 'details': [
            {'tab': 'stock', 'part': 'P1', 'location_type': 'WIP', 'quantity': '8'},
            {'tab': 'stock', 'part': 'P1', 'location_type': 'RAW', 'quantity': '2'}]})
        self.assertEqual(self.client.get('/api/material-inventory-batch/details/', {'part': 'P1'}).data['total'], '2')
        self.assertEqual(self.client.get('/api/material-inventory-batch/details/', {'part': 'P1', 'include_wip': '1'}).data['total'], '10')
        self.client.force_authenticate(get_user_model().objects.create_user('restricted', password='test'))
        self.assertEqual(self.client.get('/api/material-inventory-batch/').status_code, 403)

    def test_summary_and_details_share_capture_when_source_changes(self):
        rows = [{'tmp_part': 'P1', 'tmp_qty_oh': '1', 'tmp_qty_req_total': '0', 'tmp_qty_on_order': '0',
                 'tmp_qty_req_open': '0', 'tmp_qty_alloc': '0', 'tmp_sfty_stk': '2'}]
        differences = reconcile_bulk(rows, [{'tab': 'stock', 'part': 'P1', 'nettable': True, 'quantity': '10'},
            {'tab': 'stock', 'part': 'P1', 'nettable': False, 'quantity': '500'},
            {'tab': 'demand', 'part': 'P1', 'quantity': '4'}])
        self.assertEqual(len(differences), 2)
        self.assertEqual(rows[0]['tmp_qty_remain'], '4')
        self.assertEqual(rows[0]['tmp_qty_unalloc'], '10')

    def test_draft_requisition_is_visible_but_only_approved_supply_changes_balance(self):
        rows = [{'tmp_part': 'P1', 'tmp_qty_oh': '10', 'tmp_qty_req_total': '4',
            'tmp_qty_on_order': '0', 'tmp_qty_req_open': '0', 'tmp_qty_alloc': '0',
            'tmp_sfty_stk': '0'}]
        details = [
            {'tab': 'stock', 'part': 'P1', 'nettable': True, 'quantity': '10'},
            {'tab': 'demand', 'part': 'P1', 'quantity': '4'},
            {'tab': 'requisition', 'part': 'P1', 'quantity': '3', 'eligible_for_balance': False},
        ]
        reconcile_bulk(rows, details)
        self.assertEqual(rows[0]['tmp_qty_req_open'], '3')
        self.assertEqual(rows[0]['tmp_qty_remain'], '6')
        details[-1]['eligible_for_balance'] = True
        reconcile_bulk(rows, details)
        self.assertEqual(rows[0]['tmp_qty_req_open'], '3')
        self.assertEqual(rows[0]['tmp_qty_remain'], '9')
